#!/usr/bin/env python3
"""Menghitung kesepakatan peninjau manusia dan penyaring AI untuk Cek 1.

Dijalankan SETELAH kolom ``keputusan_manusia`` di lembar sampel diisi. Skrip ini:

1. membaca lembar isian (``sampel_judul*.csv``/``.xlsx`` dan
   ``sampel_abstrak*.csv``/``.xlsx``) dan kunci AI di ``verifikasi/.kunci/``;
2. menulis salinan terbuka (kolom ``keputusan_ai`` terisi) ke
   ``verifikasi/terbuka/`` tanpa mengubah lembar isian asli;
3. menghitung per tahap: n, persentase kesepakatan, kappa Cohen (implementasi
   sendiri), matriks kebingungan; tahap abstrak dihitung pada kode penuh
   (C1..C5, T, R, X-E1..X-E7), kode tanpa alasan eksklusi (X digabung), dan
   biner masuk/keluar;
4. menandai rekaman yang dikeluarkan AI tetapi dinilai manusia C1 atau C3
   (tahap judul: ``L-C1``/``L-C3``) — aturan lulus Cek 1;
5. menulis ``verifikasi/kesepakatan.md`` dan ``verifikasi/ketidaksepakatan.csv``
   (semua ketidaksepakatan, untuk diputuskan Fatma). Isian
   ``keputusan_final_fatma`` yang sudah ada di ``ketidaksepakatan.csv`` dipertahankan.

Pemakaian:
    python3 tools/scopus/verifikasi_kappa.py              # --sumber auto
    python3 tools/scopus/verifikasi_kappa.py --sumber xlsx
"""
import argparse
import csv
import datetime as dt
import glob
import io
import math
import re
import sys
from collections import Counter
from pathlib import Path

ROOT = Path(__file__).resolve().parents[2]
VER = ROOT / "literature/scopus-2026-09/verifikasi"
KUNCI = VER / ".kunci"
TERBUKA = VER / "terbuka"

KODE_ABSTRAK = ["C1", "C2", "C3", "C4", "C5", "T", "R",
                "X-E1", "X-E2", "X-E3", "X-E4", "X-E5", "X-E6", "X-E7"]
KODE_ABSTRAK_8 = ["C1", "C2", "C3", "C4", "C5", "T", "R", "X"]
KODE_JUDUL = ["L", "L-C1", "L-C3", "X"]
KRITIS_ABSTRAK = {"C1", "C3"}
KRITIS_JUDUL = {"L-C1", "L-C3"}

KOLOM_TIDAK_SEPAKAT = ["tahap", "berkas", "no", "idx", "eid", "key", "judul",
                       "keputusan_manusia", "keputusan_ai", "kode_ai_mentah",
                       "kritis", "catatan", "keputusan_final_fatma", "alasan_fatma"]


# ------------------------------------------------------------- baca & tulis
def baca_csv(path):
    """Baca CSV UTF-8 (dengan/tanpa BOM); toleran terhadap simpanan ulang Excel
    (pengodean cp1252 dan pemisah titik koma pada lokal Indonesia)."""
    data = Path(path).read_bytes()
    try:
        teks = data.decode("utf-8-sig")
    except UnicodeDecodeError:
        teks = data.decode("cp1252")
    kepala = teks.split("\n", 1)[0]
    pemisah = ";" if kepala.count(";") > kepala.count(",") else ","
    return list(csv.DictReader(io.StringIO(teks, newline=""), delimiter=pemisah))


def tulis_csv(path, kolom, baris):
    path.parent.mkdir(parents=True, exist_ok=True)
    with open(path, "w", encoding="utf-8-sig", newline="") as f:
        w = csv.DictWriter(f, fieldnames=kolom, extrasaction="ignore", lineterminator="\n")
        w.writeheader()
        w.writerows(baris)


def baca_xlsx(path):
    from openpyxl import load_workbook
    ws = load_workbook(path, read_only=True, data_only=True)["sampel"]
    it = ws.iter_rows(values_only=True)
    kolom = [str(c) for c in next(it)]
    out = []
    for r in it:
        if r is None or all(v is None for v in r):
            continue
        out.append({k: ("" if v is None else str(v)) for k, v in zip(kolom, r)})
    return out


def jumlah_isian(baris):
    return sum(1 for r in baris if (r.get("keputusan_manusia") or "").strip())


def pilih_sumber(csv_path, sumber):
    xl = csv_path.with_suffix(".xlsx")
    b_csv = baca_csv(csv_path)
    if sumber == "csv":
        return b_csv, csv_path.name
    if sumber == "xlsx":
        if not xl.exists():
            sys.exit(f"{xl.name} tidak ada")
        return baca_xlsx(xl), xl.name
    # auto
    b_xl = []
    if xl.exists():
        try:
            b_xl = baca_xlsx(xl)
        except ImportError:
            print(f"openpyxl tidak tersedia; {xl.name} tidak dibaca")
    n_csv, n_xl = jumlah_isian(b_csv), jumlah_isian(b_xl)
    if n_csv and n_xl:
        sys.exit(f"{csv_path.name} dan {xl.name} sama-sama berisi keputusan; pilih --sumber csv atau xlsx")
    if n_xl:
        return b_xl, xl.name
    return b_csv, csv_path.name


# ------------------------------------------------------------- normalisasi
def norm_judul(v):
    v = re.sub(r"[\s_]+", "-", (v or "").strip().upper())
    v = {"LANJUT": "L", "I": "L", "EKSKLUSI": "X", "E": "X", "YES": "L", "Y": "L", "YA": "L",
         "NO": "X", "N": "X", "TIDAK": "X", "YES-C1": "L-C1", "YES-C3": "L-C3"}.get(v, v)
    v = re.sub(r"^L-?(C[13])$", r"L-\1", v)
    return v


def norm_abstrak(v):
    v = (v or "").strip().upper().replace("_", "-")
    m = re.fullmatch(r"X\s*-?\s*(E[1-7])", v)
    if m:
        return f"X-{m.group(1)}"
    return v


def biner_judul(v):
    return "lanjut" if v.startswith("L") else "eksklusi"


def biner_abstrak(v):
    return "eksklusi" if v.startswith("X") else "masuk"


def delapan(v):
    return "X" if v.startswith("X") else v


# ------------------------------------------------------------- statistik
def kappa(pasangan, label=None):
    """Kappa Cohen untuk daftar (manusia, ai). Mengembalikan dict statistik."""
    n = len(pasangan)
    if n == 0:
        return {"n": 0}
    label = label or sorted({x for p in pasangan for x in p})
    po = sum(1 for a, b in pasangan if a == b) / n
    cm = Counter(a for a, _ in pasangan)
    ca = Counter(b for _, b in pasangan)
    pe = sum(cm[k] * ca[k] for k in label) / (n * n)
    if pe >= 1:
        k, se = float("nan"), float("nan")
    else:
        k = (po - pe) / (1 - pe)
        # galat baku pendekatan sederhana (Cohen 1960); selang 95% hanya indikatif
        se = math.sqrt(po * (1 - po) / (n * (1 - pe) ** 2))
    return {"n": n, "po": po, "pe": pe, "kappa": k, "se": se,
            "bawah": k - 1.96 * se, "atas": k + 1.96 * se}


def matriks(pasangan, label):
    c = Counter(pasangan)
    return [[c[(a, b)] for b in label] for a in label]


def fmt(x, d=3):
    if x is None or (isinstance(x, float) and math.isnan(x)):
        return "tak terdefinisi"
    return f"{x:.{d}f}".replace(".", ",")


def pct(x):
    return fmt(100 * x, 1) + "%"


def tabel_matriks(pasangan, label, judul):
    lab = [l for l in label if any(l in p for p in pasangan)] or label
    m = matriks(pasangan, lab)
    out = [f"**{judul}** (baris = manusia, kolom = AI)", "",
           "| manusia \\ AI | " + " | ".join(lab) + " | jumlah |",
           "|---|" + "---:|" * (len(lab) + 1)]
    for l, baris in zip(lab, m):
        out.append(f"| {l} | " + " | ".join(str(x) for x in baris) + f" | {sum(baris)} |")
    kol = [sum(m[i][j] for i in range(len(lab))) for j in range(len(lab))]
    out.append("| jumlah | " + " | ".join(str(x) for x in kol) + f" | {sum(kol)} |")
    return out + [""]


def baris_stat(nama, s):
    if not s.get("n"):
        return f"| {nama} | 0 | – | – | – |"
    return (f"| {nama} | {s['n']} | {pct(s['po'])} | {fmt(s['kappa'])} | "
            f"{fmt(s['bawah'])} – {fmt(s['atas'])} |")


# ------------------------------------------------------------- per tahap
def muat_tahap(tahap, sumber):
    lembar = sorted(glob.glob(str(VER / f"sampel_{tahap}*.csv")))
    kunci = {}
    for fn in sorted(glob.glob(str(KUNCI / f"kunci_{tahap}*.csv"))):
        for r in baca_csv(fn):
            kunci[r["idx"]] = r
    data, masalah, belum, terbuka = [], [], [], []
    sah = KODE_JUDUL if tahap == "judul" else KODE_ABSTRAK
    norm = norm_judul if tahap == "judul" else norm_abstrak
    for fn in lembar:
        p = Path(fn)
        baris, nama = pilih_sumber(p, sumber)
        for r in baris:
            k = kunci.get(r["idx"])
            if k is None or k["eid"] != r["eid"]:
                sys.exit(f"{nama} idx {r['idx']}: tidak ada di kunci atau EID tidak cocok")
            h_mentah = (r.get("keputusan_manusia") or "").strip()
            h = norm(h_mentah)
            buka = dict(r)
            buka["keputusan_ai"] = k["keputusan_ai"]
            terbuka.append((p.name, buka))
            if not h_mentah:
                belum.append((nama, r))
                continue
            if h not in sah:
                masalah.append(f"{nama} no {r['no']} (idx {r['idx']}): '{h_mentah}'")
                continue
            data.append({"berkas": nama, "r": r, "h": h, "ai": k["keputusan_ai"],
                         "mentah": k.get("kode_ai_mentah", ""), "k": k})
    return lembar, data, masalah, belum, terbuka


def tulis_terbuka(terbuka):
    per = {}
    for nama, b in terbuka:
        per.setdefault(nama, []).append(b)
    for nama, baris in per.items():
        tulis_csv(TERBUKA / nama.replace(".csv", "_terbuka.csv"), list(baris[0].keys()), baris)


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--sumber", choices=["auto", "csv", "xlsx"], default="auto")
    ap.add_argument("--izinkan-belum-lengkap", action="store_true",
                    help="tetap hitung walau ada baris kosong (membuka kunci AI lebih awal; hindari)")
    a = ap.parse_args()

    if not KUNCI.exists():
        sys.exit("folder .kunci tidak ada; jalankan verifikasi_sampel.py lebih dulu")

    # Penjaga kebutaan: jangan membuka keputusan AI sebelum semua baris terisi.
    kosong = {t: len(muat_tahap(t, a.sumber)[3]) for t in ("judul", "abstrak")}
    if any(kosong.values()) and not a.izinkan_belum_lengkap:
        sys.exit(f"Masih ada baris keputusan_manusia yang kosong (judul {kosong['judul']}, "
                 f"abstrak {kosong['abstrak']}). Tidak ada berkas yang ditulis agar peninjau "
                 "tetap buta. Selesaikan pengisian, atau pakai --izinkan-belum-lengkap.")

    md = ["# Kesepakatan Cek 1: peninjau manusia vs penyaring AI", "",
          f"Dihitung {dt.datetime.now(dt.timezone.utc).strftime('%Y-%m-%d %H:%M UTC')} oleh "
          "`tools/scopus/verifikasi_kappa.py`. Baris = keputusan manusia, kolom = keputusan AI. "
          "Selang 95% kappa memakai galat baku pendekatan sederhana dan hanya indikatif.", ""]
    semua_tidak_sepakat, kritis_total, belum_total, masalah_total = [], 0, 0, []
    ringkas = ["| Ukuran | n | Kesepakatan | Kappa | Selang 95% |", "|---|---:|---:|---:|---|"]
    detail = []

    for tahap in ("judul", "abstrak"):
        lembar, data, masalah, belum, terbuka = muat_tahap(tahap, a.sumber)
        if not lembar:
            detail += [f"## Tahap {tahap}", "", "Belum ada lembar sampel.", ""]
            continue
        tulis_terbuka(terbuka)
        masalah_total += masalah
        belum_total += len(belum)
        n_lembar = len(data) + len(belum) + len(masalah)
        detail += [f"## Tahap {tahap}", "",
                   f"Lembar: {', '.join(Path(x).name for x in lembar)}. "
                   f"Baris {n_lembar}; terisi sah {len(data)}; belum diisi {len(belum)}; "
                   f"nilai tidak sah {len(masalah)}.", ""]
        if tahap == "judul":
            pb = [(biner_judul(d["h"]), biner_judul(d["ai"])) for d in data]
            s = kappa(pb, ["lanjut", "eksklusi"])
            ringkas.append(baris_stat("Judul, biner lanjut/eksklusi", s))
            if s.get("n"):
                detail.append(f"PABAK (kappa disesuaikan prevalensi dan bias) = {fmt(2 * s['po'] - 1)}.")
                detail.append("")
            detail += tabel_matriks(pb, ["lanjut", "eksklusi"], "Biner")
            kritis = lambda d: d["ai"] == "X" and d["h"] in KRITIS_JUDUL
            beda = lambda d: biner_judul(d["h"]) != biner_judul(d["ai"])
        else:
            pf = [(d["h"], d["ai"]) for d in data]
            p8 = [(delapan(d["h"]), delapan(d["ai"])) for d in data]
            pb = [(biner_abstrak(d["h"]), biner_abstrak(d["ai"])) for d in data]
            s_f = kappa(pf, KODE_ABSTRAK)
            s_8 = kappa(p8, KODE_ABSTRAK_8)
            s_b = kappa(pb, ["masuk", "eksklusi"])
            ringkas += [baris_stat("Abstrak, kode penuh (14 kategori)", s_f),
                        baris_stat("Abstrak, kode dengan X digabung (8)", s_8),
                        baris_stat("Abstrak, biner masuk/eksklusi", s_b)]
            if s_b.get("n"):
                detail.append(f"PABAK biner = {fmt(2 * s_b['po'] - 1)}.")
                detail.append("")
            detail += tabel_matriks(pb, ["masuk", "eksklusi"], "Biner")
            detail += tabel_matriks(p8, KODE_ABSTRAK_8, "Kode (X digabung)")
            px = [(d["h"], d["ai"]) for d in data if d["h"].startswith("X") and d["ai"].startswith("X")]
            if px:
                detail += tabel_matriks(px, [k for k in KODE_ABSTRAK if k.startswith("X")],
                                        "Alasan eksklusi, hanya rekaman yang dikeluarkan keduanya")
            kritis = lambda d: d["ai"].startswith("X") and d["h"] in KRITIS_ABSTRAK
            beda = lambda d: d["h"] != d["ai"]
        dk = [d for d in data if kritis(d)]
        kritis_total += len(dk)
        detail.append(f"Rekaman kritis (AI mengeluarkan, manusia menilai C1/C3): **{len(dk)}**")
        detail.append("")
        for d in dk:
            detail.append(f"- {d['berkas']} no {d['r']['no']}, idx {d['r']['idx']}: manusia {d['h']}, "
                          f"AI {d['ai']} — {d['r']['judul'][:110]}")
        if dk:
            detail.append("")
        if masalah:
            detail += ["Nilai tidak sah (tidak dihitung):", ""] + [f"- {m}" for m in masalah] + [""]
        for d in data:
            if beda(d):
                semua_tidak_sepakat.append({
                    "tahap": tahap, "berkas": d["berkas"], "no": d["r"]["no"],
                    "idx": d["r"]["idx"], "eid": d["r"]["eid"], "key": d["r"].get("key", ""),
                    "judul": d["r"]["judul"], "keputusan_manusia": d["h"],
                    "keputusan_ai": d["ai"], "kode_ai_mentah": d["mentah"],
                    "kritis": "ya" if kritis(d) else "tidak",
                    "catatan": d["r"].get("catatan", ""),
                    "keputusan_final_fatma": "", "alasan_fatma": ""})

    # pertahankan keputusan Fatma yang sudah ada
    path_ts = VER / "ketidaksepakatan.csv"
    if path_ts.exists():
        lama = {(r["tahap"], r["idx"]): r for r in baca_csv(path_ts)}
        for b in semua_tidak_sepakat:
            r = lama.get((b["tahap"], b["idx"]))
            if r:
                b["keputusan_final_fatma"] = r.get("keputusan_final_fatma", "")
                b["alasan_fatma"] = r.get("alasan_fatma", "")
        hilang = set(lama) - {(b["tahap"], b["idx"]) for b in semua_tidak_sepakat}
        hilang = [k for k in hilang if lama[k].get("keputusan_final_fatma", "").strip()]
        if hilang:
            sys.exit(f"{len(hilang)} baris ketidaksepakatan.csv berisi keputusan Fatma tetapi tidak "
                     "lagi berbeda pendapat; periksa manual sebelum menimpa: " + str(sorted(hilang)[:10]))
    semua_tidak_sepakat.sort(key=lambda b: (b["kritis"] != "ya", b["tahap"], b["berkas"], int(b["no"])))
    tulis_csv(path_ts, KOLOM_TIDAK_SEPAKAT, semua_tidak_sepakat)

    if masalah_total or belum_total:
        status = "BELUM LENGKAP"
    elif kritis_total:
        status = "GAGAL"
    else:
        status = "LULUS"
    md += [f"**Status Cek 1: {status}.** Rekaman kritis {kritis_total}; baris belum diisi "
           f"{belum_total}; nilai tidak sah {len(masalah_total)}; ketidaksepakatan "
           f"{len(semua_tidak_sepakat)} (lihat `ketidaksepakatan.csv`).", "",
           "Aturan lulus: sampel tidak boleh memuat rekaman yang dikeluarkan AI tetapi dinilai "
           "manusia C1 atau C3 (tahap judul: `L-C1`/`L-C3`). Bila gagal, ikuti langkah di "
           "`README.md` bagian 5.", "", "## Ringkasan", ""] + ringkas + [""] + detail
    (VER / "kesepakatan.md").write_text("\n".join(md) + "\n", encoding="utf-8")
    print("\n".join(ringkas))
    print(f"Status: {status}; kritis {kritis_total}; belum diisi {belum_total}; "
          f"tidak sah {len(masalah_total)}; ketidaksepakatan {len(semua_tidak_sepakat)}")
    if masalah_total:
        sys.exit(2)


if __name__ == "__main__":
    main()
