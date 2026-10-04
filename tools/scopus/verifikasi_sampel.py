#!/usr/bin/env python3
"""Menarik sampel acak buta untuk Cek 1 (verifikasi penyaringan oleh manusia).

Sampel ditarik dengan benih (*seed*) tetap sehingga dapat diulang persis:

- 300 rekaman dari populasi tahap judul (5.723 rekaman: semua baris
  ``tahap_judul.csv`` kecuali ``X0`` = dikeluarkan menurut tipe dokumen);
- 100 rekaman dari populasi tahap kelayakan/abstrak (1.124 kandidat di
  ``kandidat_abstrak.csv``).

Lembar isian bersifat buta: kolom ``keputusan_ai`` sengaja dikosongkan.
Keputusan AI disimpan terpisah di ``verifikasi/cek-1-sampel-buta/.kunci/`` dan baru digabung oleh
``verifikasi_kappa.py`` setelah peninjau manusia selesai mengisi.

Keluaran (folder ``literature/scopus-2026-09/verifikasi/cek-1-sampel-buta/``):
    sampel_judul.csv, sampel_abstrak.csv       lembar isian buta (CSV, UTF-8 BOM)
    sampel_judul.xlsx, sampel_abstrak.xlsx     versi Excel dengan daftar pilihan
    .kunci/kunci_judul.csv, .kunci/kunci_abstrak.csv   keputusan AI per idx
    .kunci/riwayat_sampel.csv                  catatan setiap penarikan

Pemakaian:
    python3 tools/scopus/verifikasi_sampel.py                 # sampel awal
    python3 tools/scopus/verifikasi_sampel.py --tambah 300 --benih 20261006
        # sampel tambahan tahap judul, tanpa idx yang sudah pernah disampel

Aturan keputusan AI tahap kelayakan: berkas ``abstrak_*.txt`` dibaca urut nama
(``sorted``: 00..08, lalu B00..B03, lalu T00); bila satu idx muncul lebih dari
sekali, baris terakhir yang berlaku. Aturan ini sama dengan
``kode_bukti.py`` dan ``gambar_tinjauan.py`` sehingga kunci identik dengan
keputusan yang dipakai naskah.
"""
import argparse
import csv
import datetime as dt
import glob
import io
import json
import random
import re
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[2]
KORPUS = ROOT / "literature/scopus-2026-09"
TOPIK = KORPUS / "topik"
PEN = TOPIK / "penyaringan"
VER = KORPUS / "verifikasi" / "cek-1-sampel-buta"
PUTARAN_1 = KORPUS / "verifikasi" / "hasil-kerja-ai" / "putaran_1"
KUNCI = VER / ".kunci"

BENIH_BAWAAN = 20260929
N_JUDUL = 300
N_ABSTRAK = 100
POP_JUDUL_DIHARAPKAN = 5723
POP_ABSTRAK_DIHARAPKAN = 1124

KODE = re.compile(r"^(\d+)\s+(C[1-5]|T|R|M|X)(?:\s+(E\d))?(.*)$")
LANJUT_JUDUL = {"I", "M", "R", "T"}          # I, M (ragu), R, T -> lanjut ke abstrak

PILIHAN_JUDUL = ["Yes", "Yes-C1", "Yes-C3", "No"]
PILIHAN_ABSTRAK = ["C1", "C2", "C3", "C4", "C5", "T", "R",
                   "X-E1", "X-E2", "X-E3", "X-E4", "X-E5", "X-E6", "X-E7"]

KOLOM_JUDUL = ["no", "idx", "eid", "judul", "tahun", "sumber",
               "keputusan_manusia", "catatan", "keputusan_ai", "keputusan_final_fatma"]
KOLOM_ABSTRAK = ["no", "idx", "eid", "key", "judul", "tahun", "sumber", "abstrak",
                 "keputusan_manusia", "catatan", "keputusan_ai", "keputusan_final_fatma"]
KOLOM_KUNCI_JUDUL = ["idx", "eid", "kode_ai_mentah", "keputusan_ai"]
KOLOM_KUNCI_ABSTRAK = ["idx", "eid", "key", "kode_ai_mentah", "keputusan_ai",
                       "berkas_sumber", "dasar_ai"]
KOLOM_RIWAYAT = ["waktu_utc", "tahap", "berkas", "benih", "n", "populasi",
                 "dikecualikan_sudah_disampel"]

TANPA_ABSTRAK = ("[Abstrak tidak tersedia di korpus. Putuskan dari judul, sumber, "
                 "dan tahun; bila ragu, catat di kolom catatan.]")


def baca_csv(path):
    """Baca CSV UTF-8 (dengan/tanpa BOM); toleran terhadap simpanan ulang Excel
    (pengodean cp1252 dan pemisah titik koma pada lokal Indonesia)."""
    data = Path(path).read_bytes()
    try:
        teks = data.decode("utf-8-sig")
    except UnicodeDecodeError:
        teks = data.decode("cp1252")
    kepala = teks.split("\n", 1)[0]
    pemisah = ";" if kepala.count(";") > kepala.count(",") else ","
    return list(csv.DictReader(io.StringIO(teks, newline=""), delimiter=pemisah))


def tulis_csv(path, kolom, baris):
    path.parent.mkdir(parents=True, exist_ok=True)
    with open(path, "w", encoding="utf-8-sig", newline="") as f:
        w = csv.DictWriter(f, fieldnames=kolom, lineterminator="\n")  # LF, sesuai .gitattributes
        w.writeheader()
        w.writerows(baris)


# ------------------------------------------------------------------ populasi
def populasi_judul():
    tj = baca_csv(PEN / "tahap_judul.csv")
    rek = {r["idx"]: r for r in baca_csv(TOPIK / "records_all.csv")}
    pop = {}
    for r in tj:
        if r["keputusan_judul"] == "X0":          # dikeluarkan menurut tipe dokumen
            continue
        m = rek[r["idx"]]
        if m["eid"] != r["eid"]:
            sys.exit(f"idx {r['idx']}: EID tahap_judul.csv dan records_all.csv berbeda")
        kode = r["keputusan_judul"]
        pop[int(r["idx"])] = {
            "idx": r["idx"], "eid": r["eid"], "judul": m["title"],
            "tahun": m["year"], "sumber": m["source"],
            "kode_ai_mentah": kode,
            "keputusan_ai": "L" if kode in LANJUT_JUDUL else "X",
        }
    return pop


def keputusan_abstrak(folder=None):
    """idx -> (kode, alasan, berkas, dasar). Urut nama berkas, baris terakhir berlaku."""
    dec, riwayat = {}, {}
    for fn in sorted(glob.glob(str((folder or PEN) / "abstrak_*.txt"))):
        nama = Path(fn).name
        with open(fn, encoding="utf-8") as f:
            for line in f:
                m = KODE.match(line.strip())
                if not m:
                    continue
                i = int(m.group(1))
                dasar = "judul" if "(judul)" in line or nama == "abstrak_T00.txt" else "abstrak"
                nilai = (m.group(2), m.group(3) or "", nama, dasar)
                riwayat.setdefault(i, []).append(nilai)
                dec[i] = nilai
    ganda = {i: v for i, v in riwayat.items() if len(v) > 1}
    return dec, ganda


def populasi_abstrak():
    kand = baca_csv(PEN / "kandidat_abstrak.csv")
    enr = {}
    with open(TOPIK / "enrich.jsonl", encoding="utf-8") as f:
        for line in f:
            e = json.loads(line)
            enr[e["eid"]] = e
    dec, ganda = keputusan_abstrak()
    pop = {}
    for r in kand:
        i = int(r["idx"])
        if i not in dec:
            sys.exit(f"idx {i} tidak punya keputusan di abstrak_*.txt")
        kode, alasan, berkas, dasar = dec[i]
        if kode == "M":
            sys.exit(f"idx {i} masih berkode M (ragu) di {berkas}")
        if kode == "X" and not alasan:
            sys.exit(f"idx {i} berkode X tanpa alasan E1–E7 di {berkas}")
        abstrak = (enr.get(r["eid"], {}).get("abstract") or "").strip()
        pop[i] = {
            "idx": r["idx"], "eid": r["eid"], "key": r["key"], "judul": r["title"],
            "tahun": r["year"], "sumber": r["source"],
            "abstrak": abstrak or TANPA_ABSTRAK,
            "kode_ai_mentah": f"{kode} {alasan}".strip(),
            "keputusan_ai": f"X-{alasan}" if kode == "X" else kode,
            "berkas_sumber": berkas, "dasar_ai": dasar,
        }
    lebih = set(dec) - set(pop)
    if lebih:
        print(f"PERINGATAN: {len(lebih)} idx di abstrak_*.txt bukan kandidat_abstrak.csv")
    return pop, ganda


def abstrak_bersarang():
    """Sampel abstrak yang diambil dari sampel judul: semua rekaman sampel judul yang oleh AI
    putaran 1 dinyatakan lanjut ke abstrak (kunci_judul.csv), dengan keputusan abstrak AI
    putaran 1 (salinan ai2/putaran_1). Urutan mengikuti sampel judul."""
    sampel = [int(r["idx"]) for r in baca_csv(VER / "sampel_judul.csv")]
    lanjut = {int(r["idx"]) for r in baca_csv(KUNCI / "kunci_judul.csv") if r["keputusan_ai"] == "L"}
    pilih = [i for i in sampel if i in lanjut]
    kand = {int(r["idx"]): r for r in baca_csv(PEN / "kandidat_abstrak.csv")}
    enr = {}
    with open(TOPIK / "enrich.jsonl", encoding="utf-8") as f:
        for line in f:
            e = json.loads(line)
            enr[e["eid"]] = e
    dec, _ = keputusan_abstrak(PUTARAN_1)
    pop = {}
    for i in pilih:
        if i not in kand or i not in dec:
            sys.exit(f"idx {i} lanjut di putaran 1 tetapi tidak ada di kandidat_abstrak.csv atau keputusan putaran 1")
        r = kand[i]
        kode, alasan, berkas, dasar = dec[i]
        abstrak = (enr.get(r["eid"], {}).get("abstract") or "").strip()
        pop[i] = {
            "idx": r["idx"], "eid": r["eid"], "key": r["key"], "judul": r["title"],
            "tahun": r["year"], "sumber": r["source"], "abstrak": abstrak or TANPA_ABSTRAK,
            "kode_ai_mentah": f"{kode} {alasan}".strip(),
            "keputusan_ai": f"X-{alasan}" if kode == "X" else kode,
            "berkas_sumber": berkas, "dasar_ai": dasar,
        }
    return pop, pilih


# --------------------------------------------------------------------- xlsx
def tulis_xlsx(path, kolom, baris, pilihan, tahap, benih):
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter
        from openpyxl.worksheet.datavalidation import DataValidation
    except ImportError:
        print("openpyxl tidak tersedia; versi .xlsx dilewati")
        return False
    wb = Workbook()
    ws = wb.active
    ws.title = "sampel"
    ws.append(kolom)
    for b in baris:
        ws.append([b.get(k, "") for k in kolom])
    n = len(baris) + 1
    lebar = {"no": 5, "idx": 7, "eid": 21, "key": 22, "judul": 60, "tahun": 7,
             "sumber": 30, "abstrak": 90, "keputusan_manusia": 18, "catatan": 35,
             "keputusan_ai": 13, "keputusan_final_fatma": 18}
    for j, k in enumerate(kolom, 1):
        ws.column_dimensions[get_column_letter(j)].width = lebar.get(k, 12)
        ws.cell(1, j).font = Font(bold=True)
    isi = PatternFill("solid", fgColor="FFF2CC")
    for k in ("keputusan_manusia", "catatan", "keputusan_final_fatma"):
        j = kolom.index(k) + 1
        for r in range(2, n + 1):
            ws.cell(r, j).fill = isi
    for k in ("judul", "abstrak", "catatan", "sumber"):
        if k in kolom:
            j = kolom.index(k) + 1
            for r in range(2, n + 1):
                ws.cell(r, j).alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "A2"
    daftar = '"' + ",".join(pilihan) + '"'
    for k in ("keputusan_manusia", "keputusan_final_fatma"):
        dv = DataValidation(type="list", formula1=daftar, allow_blank=True,
                            showErrorMessage=True, errorTitle="Nilai tidak dikenal",
                            error="Pilih salah satu nilai dari daftar.")
        huruf = get_column_letter(kolom.index(k) + 1)
        dv.add(f"{huruf}2:{huruf}{n}")
        ws.add_data_validation(dv)
    p = wb.create_sheet("petunjuk")
    p.column_dimensions["A"].width = 14
    p.column_dimensions["B"].width = 100
    p.append(["Cek 1", f"Sampel buta tahap {tahap}, benih {benih}. Kolom keputusan_ai "
                        "sengaja kosong. Lihat verifikasi/cek-1-sampel-buta/PANDUAN.md."])
    if tahap == "judul":
        teks = [("Yes", "Lanjut ke abstrak (judul mungkin relevan; bila ragu, pilih Yes)"),
                ("Yes-C1", "Lanjut; judul menunjukkan beberapa pengamatan buah yang sama (video, multipandang, pindai ulang)"),
                ("Yes-C3", "Lanjut; judul menyangkut pencitraan TBS kelapa sawit"),
                ("No", "Eksklusi pada tahap judul")]
    else:
        teks = [("C1", "Menggabungkan beberapa pengamatan buah yang sama (video, beberapa pandang, pindaian berulang)"),
                ("C2", "Pencacahan atau estimasi hasil dari satu pandang"),
                ("C3", "Pencitraan TBS kelapa sawit, termasuk grading di pabrik dan brondolan"),
                ("C4", "Atribut kelas disertai pencacahan dari citra tunggal"),
                ("C5", "Deteksi, lokalisasi, atau pengukuran buah dengan depth, 3D, atau modalitas non-RGB"),
                ("T", "Metode yang dapat dipindahkan dari luar pertanian"),
                ("R", "Tinjauan terdahulu"),
                ("X-E1", "Eksklusi: bukan buah pada tanaman"),
                ("X-E2", "Eksklusi: hanya pascapanen atau laboratorium"),
                ("X-E3", "Eksklusi: model hasil tanpa deteksi tingkat buah"),
                ("X-E4", "Eksklusi: hanya pemetikan atau manipulasi"),
                ("X-E5", "Eksklusi: buah dideteksi tetapi tidak dicacah/dievaluasi lebih lanjut"),
                ("X-E6", "Eksklusi: sensor non-citra"),
                ("X-E7", "Eksklusi: bahasa selain Inggris")]
    for a, b in teks:
        p.append([a, b])
    wb.save(path)
    return True


# --------------------------------------------------------------------- main
def sudah_disampel(tahap):
    idx = set()
    for fn in glob.glob(str(VER / f"sampel_{tahap}*.csv")):
        idx |= {int(r["idx"]) for r in baca_csv(fn)}
    return idx


def benih_terpakai():
    path = KUNCI / "riwayat_sampel.csv"
    return {int(r["benih"]) for r in baca_csv(path)} if path.exists() else set()


def catat_riwayat(tahap, berkas, benih, n, pop, kecuali):
    path = KUNCI / "riwayat_sampel.csv"
    baru = not path.exists()
    KUNCI.mkdir(parents=True, exist_ok=True)
    with open(path, "a", encoding="utf-8-sig" if baru else "utf-8", newline="") as f:
        w = csv.DictWriter(f, fieldnames=KOLOM_RIWAYAT, lineterminator="\n")
        if baru:
            w.writeheader()
        w.writerow({"waktu_utc": dt.datetime.now(dt.timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ"),
                    "tahap": tahap, "berkas": berkas, "benih": benih, "n": n,
                    "populasi": pop, "dikecualikan_sudah_disampel": kecuali})


def ada_isian(path):
    """True bila lembar (CSV atau XLSX) sudah memuat keputusan manusia/Fatma."""
    if not path.exists():
        return False
    if path.suffix == ".xlsx":
        from openpyxl import load_workbook
        wb = load_workbook(path, read_only=True)
        try:
            it = wb["sampel"].iter_rows(values_only=True)
            kolom = list(next(it))
            j = [kolom.index(k) for k in ("keputusan_manusia", "catatan", "keputusan_final_fatma")]
            return any(r and any(r[x] not in (None, "") for x in j) for r in it)
        finally:
            wb.close()          # Windows mengunci berkas selama masih terbuka
    return any((r.get("keputusan_manusia") or "").strip() or (r.get("catatan") or "").strip() or
               (r.get("keputusan_final_fatma") or "").strip() for r in baca_csv(path))


def tarik(pop, n, rng, kecuali=frozenset()):
    kandidat = sorted(i for i in pop if i not in kecuali)
    if n > len(kandidat):
        sys.exit(f"n={n} melebihi sisa populasi {len(kandidat)}")
    return rng.sample(kandidat, n)   # urutan acak, bukan urut idx


def buat(tahap, pop, pilih, akhiran, benih):
    if tahap == "judul":
        kolom, kolom_kunci, pilihan = KOLOM_JUDUL, KOLOM_KUNCI_JUDUL, PILIHAN_JUDUL
    else:
        kolom, kolom_kunci, pilihan = KOLOM_ABSTRAK, KOLOM_KUNCI_ABSTRAK, PILIHAN_ABSTRAK
    lembar, kunci = [], []
    for no, i in enumerate(pilih, 1):
        r = pop[i]
        b = {k: r.get(k, "") for k in kolom}
        b["no"] = no
        b["keputusan_manusia"] = b["catatan"] = ""
        b["keputusan_ai"] = ""                     # BUTA: sengaja kosong
        b["keputusan_final_fatma"] = ""
        lembar.append(b)
        kunci.append({k: r.get(k, "") for k in kolom_kunci})
    nama = f"sampel_{tahap}{akhiran}"
    tulis_csv(VER / f"{nama}.csv", kolom, lembar)
    tulis_csv(KUNCI / f"kunci_{tahap}{akhiran}.csv", kolom_kunci, kunci)
    xl = tulis_xlsx(VER / f"{nama}.xlsx", kolom, lembar, pilihan, tahap, benih)
    return nama, xl


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--benih", type=int, default=None,
                    help=f"benih acak (bawaan sampel awal {BENIH_BAWAAN}; wajib diisi untuk --tambah)")
    ap.add_argument("--tambah", type=int, default=None,
                    help="tarik N rekaman tambahan, tanpa idx yang sudah pernah disampel")
    ap.add_argument("--tahap", choices=["judul", "abstrak"], default="judul",
                    help="tahap untuk --tambah (bawaan judul)")
    ap.add_argument("--abstrak-dari-judul", action="store_true",
                    help="ganti sampel abstrak (yang BELUM diisi) dengan semua rekaman sampel judul "
                         "yang lolos ke abstrak menurut AI putaran 1")
    ap.add_argument("--timpa", action="store_true",
                    help="izinkan menimpa sampel awal yang BELUM diisi")
    a = ap.parse_args()

    if a.abstrak_dari_judul:
        akhir = [VER / f"sampel_abstrak.{e}" for e in ("csv", "xlsx")]
        if any(ada_isian(x) for x in akhir if x.exists()):
            sys.exit("sampel_abstrak sudah berisi keputusan manusia; tidak akan ditimpa.")
        cad = VER / "cadangan_sampel_abstrak_terpisah"
        cad.mkdir(exist_ok=True)
        for x in akhir + [KUNCI / "kunci_abstrak.csv"]:
            if x.exists():
                x.replace(cad / x.name)
        pop, pilih = abstrak_bersarang()
        benih = BENIH_BAWAAN if a.benih is None else a.benih
        nama, xl = buat("abstrak", pop, pilih, "", benih)
        catat_riwayat("abstrak", nama + ".csv (dari sampel judul)", benih, len(pilih), len(pilih), 0)
        print(f"Sampel abstrak bersarang: {len(pilih)} rekaman dari 300 judul -> {nama}.csv"
              + (" + .xlsx" if xl else "") + f"; sampel terpisah lama dipindah ke {cad.name}/")
        return

    pj = populasi_judul()
    pa, ganda = populasi_abstrak()
    print(f"Populasi tahap judul  : {len(pj)} (harapan {POP_JUDUL_DIHARAPKAN})")
    print(f"Populasi tahap abstrak: {len(pa)} (harapan {POP_ABSTRAK_DIHARAPKAN})")
    if ganda:
        beda = {i: v for i, v in ganda.items() if len({x[:2] for x in v}) > 1}
        print(f"idx ganda di abstrak_*.txt: {len(ganda)} (berbeda isi {len(beda)}); baris terakhir berlaku")
    else:
        print("idx ganda di abstrak_*.txt: 0 (aturan 'berkas terakhir berlaku' tidak terpakai)")
    for nama, pop, harap in (("judul", pj, POP_JUDUL_DIHARAPKAN), ("abstrak", pa, POP_ABSTRAK_DIHARAPKAN)):
        if len(pop) != harap:
            print(f"PERINGATAN: populasi {nama} {len(pop)} != {harap} di PROTOKOL.md")

    if a.tambah:
        if a.benih is None:
            sys.exit("--tambah wajib disertai --benih baru (misalnya tanggal penarikan, 20261006)")
        if a.benih in benih_terpakai() or a.benih == BENIH_BAWAAN:
            sys.exit(f"benih {a.benih} sudah pernah dipakai; pilih benih lain")
        pop = pj if a.tahap == "judul" else pa
        kecuali = sudah_disampel(a.tahap)
        akhiran = f"_tambah_{a.benih}"
        if (VER / f"sampel_{a.tahap}{akhiran}.csv").exists():
            sys.exit(f"sampel_{a.tahap}{akhiran}.csv sudah ada")
        pilih = tarik(pop, a.tambah, random.Random(a.benih), kecuali)
        nama, xl = buat(a.tahap, pop, pilih, akhiran, a.benih)
        catat_riwayat(a.tahap, nama + ".csv", a.benih, a.tambah, len(pop), len(kecuali))
        print(f"Tambahan: {a.tambah} rekaman tahap {a.tahap}, benih {a.benih}, "
              f"{len(kecuali)} idx yang sudah disampel dikecualikan -> {nama}.csv"
              + (" + .xlsx" if xl else ""))
        return

    benih = BENIH_BAWAAN if a.benih is None else a.benih
    for t in ("judul", "abstrak"):
        path = VER / f"sampel_{t}.csv"
        if path.exists():
            if not a.timpa:
                sys.exit(f"{path.name} sudah ada. Pakai --timpa hanya bila lembar belum diisi.")
            if ada_isian(path) or ada_isian(path.with_suffix(".xlsx")):
                sys.exit(f"{path.name} sudah berisi keputusan manusia; tidak akan ditimpa.")
    # satu generator: 300 judul ditarik lebih dulu, lalu 100 abstrak
    rng = random.Random(benih)
    pilih_j = tarik(pj, N_JUDUL, rng)
    pilih_a = tarik(pa, N_ABSTRAK, rng)
    nj, xj = buat("judul", pj, pilih_j, "", benih)
    na, xa = buat("abstrak", pa, pilih_a, "", benih)
    catat_riwayat("judul", nj + ".csv", benih, N_JUDUL, len(pj), 0)
    catat_riwayat("abstrak", na + ".csv", benih, N_ABSTRAK, len(pa), 0)
    print(f"Benih {benih}: random.Random({benih}).sample, judul lebih dulu lalu abstrak")
    print(f"Tulis {nj}.csv ({N_JUDUL}) dan {na}.csv ({N_ABSTRAK})"
          + (" + .xlsx" if xj and xa else "") + f"; kunci di {KUNCI.relative_to(ROOT)}")


if __name__ == "__main__":
    main()
