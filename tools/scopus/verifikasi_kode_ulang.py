#!/usr/bin/env python3
"""Cek 5: label kajian C1 yang dibaca ulang dari teks lengkap.

Mode:
  (tanpa opsi)       tulis verifikasi/cek-5-kode-ulang/kode_ulang_C1.csv dari
                     topik/bukti/mekanisme_C1.txt dan matriks_bukti.csv. Kolom
                     `bukti` dan `catatan` diambil dari lembar yang sudah ada.
  --sampel-fatma     tulis sampel_fatma_C1.xlsx: 20 kajian C1 acak (benih
                     20261106) untuk dikode buta Bu Fatma, dengan daftar pilihan
                     dan lembar kedua berisi arti setiap pilihan.
  --banding-fatma    bandingkan isian Bu Fatma dengan mekanisme_C1.txt.

Label C1 diubah langsung di mekanisme_C1.txt dan setiap perubahan dicatat di
verifikasi/.teknis/log_perubahan.csv (lihat verifikasi_log_manusia.py).

  python3 tools/scopus/verifikasi_kode_ulang.py
"""
import argparse
import csv
import random
import re
import sys
from pathlib import Path

from verifikasi_istilah import AKU, MEK, MET, REF, YN, arti, kode

ROOT = Path(__file__).resolve().parents[2]
KORPUS = ROOT / "literature/scopus-2026-09"
MAT = KORPUS / "topik/bukti/matriks_bukti.csv"
MANUAL = KORPUS / "topik/bukti/mekanisme_C1.txt"
CEK5 = KORPUS / "verifikasi/cek-5-kode-ulang"
LEMBAR = CEK5 / "kode_ulang_C1.csv"
SAMPEL = CEK5 / "sampel_fatma_C1.xlsx"
BENIH_FATMA, N_FATMA = 20261106, 20

KOLOM_MANUAL = ["mekanisme", "akuisisi", "per_kelas", "hasil_ringkas", "sumber_kode", "halaman",
                "referensi_hitung", "tingkat_metrik", "cara_kelas", "berubah"]
NAMA_KOLOM = {"mekanisme": "cara penghubungan", "akuisisi": "cara foto diambil", "per_kelas": "per kelas"}
KELAS = {"-": "tidak berlaku", "saat pelacakan": "kelas ditetapkan saat pelacakan",
         "sesudah pelacakan": "kelas ditetapkan sesudah pelacakan",
         "voting antarpandang": "kelas dipilih dengan suara antarpandang",
         "tidak dinyatakan": "tidak dinyatakan"}
# Kolom lembar Bu Fatma: (judul kolom, kolom mekanisme_C1.txt, pilihan kode -> arti)
ISIAN_FATMA = [("cara_penghubungan", "mekanisme", {k: v for k, v in MEK.items() if k != "DATA"}),
               ("cara_foto_diambil", "akuisisi", AKU),
               ("hitung_per_kelas_kematangan", "per_kelas", YN),
               ("acuan_hitungan", "referensi_hitung", REF),
               ("metrik_evaluasi", "tingkat_metrik", MET),
               ("cara_kelas", "cara_kelas", KELAS),
               ("dibaca_dari", "sumber_kode", {"teks lengkap": "teks lengkap", "abstrak": "abstrak"}),
               ("halaman_bukti", "halaman", None)]


def baca(path):
    with open(path, encoding="utf-8-sig", newline="") as f:
        return list(csv.DictReader(f))


def muat_manual():
    kode_c1 = {}
    for line in MANUAL.read_text(encoding="utf-8").splitlines():
        if line.startswith("#") or not line.strip():
            continue
        p = line.split("\t")
        p += [""] * (11 - len(p))
        kode_c1[p[0]] = dict(zip(KOLOM_MANUAL, p[1:11]))
    return kode_c1


def jamak(v, peta):
    return "; ".join(peta.get(p.strip(), p.strip()) for p in v.split(";") if p.strip())


def perubahan(v):
    """'ya (lama: mekanisme=M3; per_kelas=N)' -> 'ya; sebelumnya cara penghubungan = M3 (pelacakan video); ...'."""
    m = re.match(r"ya \(lama: (.*)\)$", v.strip())
    if not m:
        return {"tidak": "tidak", "": ""}.get(v.strip(), v)
    bag = []
    for b in m.group(1).split(";"):
        k, _, x = b.strip().partition("=")
        peta = {"mekanisme": MEK, "akuisisi": AKU, "per_kelas": YN}.get(k)
        bag.append(f"{NAMA_KOLOM.get(k, k)} = {arti(x, peta) if peta else x}")
    return "ya; sebelumnya " + "; ".join(bag)


PROSES = re.compile(r"isian model|pembaca buta|diserahkan peninjau|^Diisi \d+ Oktober|Kode ulang (Sonnet|Haiku|Opus)|"
                    r"Hasil ringkas tidak diubah|Belum dibaca peninjau|kode tidak berubah|^$")


def bersih(c):
    """Buang keterangan proses (pembaca model, tanggal isian); sisakan isi catatan."""
    bag = [b.strip() for b in re.split(r";\s+|\.\s+(?=[A-Z])|\.;\s*", c)]
    bag = [re.sub(r"^label buta dikoreksi agen utama setelah membaca teks lengkap: ",
                  "dikoreksi setelah membaca teks lengkap: ", b) for b in bag if not PROSES.search(b)]
    teks = "; ".join(re.sub(r"\s*\((agen utama|diperiksa agen utama)[^)]*\)", "", b) for b in bag).strip(" ;.")
    teks = re.sub(r"\s*\(\d+ Oktober 2026\)", "", teks).replace("dibaca agen utama", "dibaca")
    teks = teks.replace("cara_kelas", "cara kelas")
    return re.sub(r"\b(M[0-5])\b(?! \()", lambda m: f"{m.group(1)} ({MEK[m.group(1)]})", teks)


def bangun_lembar():
    lama = {r["idx"]: r for r in baca(LEMBAR)} if LEMBAR.exists() else {}
    manual = muat_manual()
    mat = {r["idx"]: r for r in baca(MAT)}
    baris = []
    for idx in sorted(manual, key=int):
        k, r, l = manual[idx], mat.get(idx, {}), lama.get(idx, {})
        baris.append({
            "idx": idx, "key": r.get("key", ""), "tahun": r.get("year", ""), "judul": r.get("title", ""),
            "dibaca_dari": k["sumber_kode"],
            "cara_penghubungan": arti(k["mekanisme"], MEK),
            "cara_foto_diambil": arti(k["akuisisi"], AKU),
            "hitung_per_kelas_kematangan": YN.get(k["per_kelas"], k["per_kelas"]),
            "cara_kelas": KELAS.get(k["cara_kelas"], k["cara_kelas"]),
            "acuan_hitungan": jamak(k["referensi_hitung"], REF),
            "metrik_evaluasi": jamak(k["tingkat_metrik"], MET),
            "hasil_ringkas": k["hasil_ringkas"],
            "halaman_bukti": k["halaman"],
            "bukti": l.get("bukti", l.get("draf_ai_bukti", "")),
            "berubah_saat_cek_5": perubahan(k["berubah"]),
            "diisi_oleh": "AI dari teks lengkap, disetujui peninjau 19 Oktober 2026",
            "catatan": bersih(l.get("catatan", l.get("catatan_peninjau", ""))),
        })
    for n, b in enumerate(baris, 1):
        b["no"] = n
    kolom = ["no"] + [c for c in baris[0] if c != "no"]
    with open(LEMBAR, "w", encoding="utf-8-sig", newline="") as f:
        w = csv.DictWriter(f, fieldnames=kolom)
        w.writeheader()
        w.writerows(baris)
    print(f"{len(baris)} kajian C1 -> {LEMBAR.relative_to(ROOT)}")


def sampel_fatma(idx=None):
    import openpyxl
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation

    if SAMPEL.exists():
        sys.exit(f"{SAMPEL.relative_to(ROOT)} sudah ada; hapus dahulu bila memang ingin menarik ulang.")
    mat = sorted((r for r in baca(MAT) if r["kode"] == "C1"), key=lambda r: int(r["idx"]))
    pilih = [r for r in mat if r["idx"] in idx] if idx else random.Random(BENIH_FATMA).sample(mat, N_FATMA)
    pilih.sort(key=lambda r: int(r["idx"]))
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "isian"
    kepala = ["no", "idx", "judul", "tahun", "doi", "pdf_tersedia"] + [c for c, _, _ in ISIAN_FATMA] + ["catatan"]
    ws.append(kepala)
    for n, r in enumerate(pilih, 1):
        ws.append([n, int(r["idx"]), r["title"], r["year"], r["doi"],
                   "ya" if (KORPUS / "pdf" / f"{r['key']}.pdf").exists() else "tidak"])
    kuning = PatternFill("solid", fgColor="FFF2CC")
    for j, c in enumerate(kepala, 1):
        sel = ws.cell(1, j)
        sel.font = Font(bold=True)
        if j > 6:
            sel.fill = kuning
    ref = wb.create_sheet("pilihan")
    ref.append(["kolom", "pilihan", "arti"])
    for kol, sumber, peta in ISIAN_FATMA:
        j = kepala.index(kol) + 1
        if peta is None:
            ref.append([kol, "bebas", "halaman, tabel, atau gambar tempat hasil utama ditemukan"])
            continue
        nilai = [f"{k} ({v})" if sumber in ("mekanisme", "akuisisi") else v for k, v in peta.items()]
        for k, v in zip(peta, nilai):
            ref.append([kol, v, ""])
        if kol in ("cara_penghubungan", "acuan_hitungan", "metrik_evaluasi"):
            ref.append([kol, "lebih dari satu", "tulis semuanya, dipisah ' + ', misalnya M2 + M3"])
        dv = DataValidation(type="list", formula1='"' + ",".join(nilai) + '"', showErrorMessage=False)
        ws.add_data_validation(dv)
        dv.add(f"{ws.cell(2, j).coordinate}:{ws.cell(len(pilih) + 1, j).coordinate}")
    ref.append(["catatan", "bebas", "alasan bila ragu"])
    for lembar, lebar in ((ws, {"C": 60}), (ref, {"A": 28, "B": 45, "C": 60})):
        for k, v in lebar.items():
            lembar.column_dimensions[k].width = v
        for row in lembar.iter_rows(min_row=2):
            for sel in row:
                sel.alignment = Alignment(wrap_text=True, vertical="top")
    wb.save(SAMPEL)
    print(f"{N_FATMA} kajian (benih {BENIH_FATMA}) -> {SAMPEL.relative_to(ROOT)}")


def banding_fatma():
    import openpyxl
    ws = openpyxl.load_workbook(SAMPEL).worksheets[0]
    rows = list(ws.iter_rows(values_only=True))
    kepala = list(rows[0])
    manual = muat_manual()
    banding = [(c, s, p) for c, s, p in ISIAN_FATMA if s in ("mekanisme", "akuisisi", "per_kelas",
                                                           "referensi_hitung", "tingkat_metrik", "cara_kelas")]
    beda, belum = [], 0
    for row in rows[1:]:
        r = {k: ("" if v is None else str(v)) for k, v in zip(kepala, row)}
        if not r["cara_penghubungan"].strip():
            belum += 1
            continue
        acuan = manual.get(r["idx"], {})
        d = []
        for kol, sumber, peta in banding:
            f = {kode(p, peta) for p in re.split(r" \+ |;", r[kol]) if p.strip()}
            a = {p.strip() for p in re.split(r"[+;]", acuan.get(sumber, "")) if p.strip()} or {"-"}
            f = f or {"-"}
            if f != a:
                d.append(f"{kol}: Bu Fatma {sorted(f)}, label {sorted(a)}")
        if d:
            beda.append(f"idx {r['idx']}: " + " | ".join(d))
    print("\n".join(beda))
    n = len(rows) - 1 - belum
    print(f"{len(beda)} dari {n} kajian berbeda (ambang: lebih dari 3 dari 20 -> cari penyebab dan periksa ulang)")
    if belum:
        print(f"belum diisi: {belum} kajian")


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--sampel-fatma", action="store_true")
    ap.add_argument("--banding-fatma", action="store_true")
    a = ap.parse_args()
    if a.sampel_fatma:
        sampel_fatma()
    elif a.banding_fatma:
        banding_fatma()
    else:
        bangun_lembar()


if __name__ == "__main__":
    main()
