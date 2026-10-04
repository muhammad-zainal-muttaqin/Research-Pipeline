#!/usr/bin/env python3
"""Daftar kerja teks lengkap (Pemeriksaan 4, bagian 1: permintaan PDF).

Menyusun satu baris per kajian dalam empat kelompok prioritas rencana
verifikasi teks lengkap, lalu menurunkan daftar permintaan PDF darinya.

  1  kajian di Tabel 2 dan 3 naskah (tab:acq + tab:assoc)
  2  kajian inti C1 yang dikodekan dari judul saja (dasar_keputusan = judul)
  3  sisa kajian C1
  4  kajian sawit C3 yang dibahas rinci: pencacahan (tugas HIT) dan
     multipandang (kolom multipandang = Y) di topik/bukti/kode_C3.txt

Keluaran (literature/scopus-2026-09/verifikasi/cek-4-teks-lengkap/ dan unduhan/):
  verifikasi/cek-4-teks-lengkap/teks_lengkap.csv       daftar kerja (dibaca dan ditulis ulang)
  verifikasi/cek-4-teks-lengkap/PERMINTAAN-PDF.md      daftar kajian tanpa PDF
  verifikasi/cek-4-teks-lengkap/PERMINTAAN-PDF-BU-FATMA.xlsx  daftar yang sama (Excel)
  unduhan/PRIORITAS-UNDUH-MANUAL.md daftar centang unduh manual

Aman dijalankan ulang: kolom pdf_ada dihitung ulang dari folder pdf/,
sedangkan kolom isian tangan (status, tanggal_minta, sumber_pdf, catatan,
email_penulis_korespondensi) yang sudah berisi dipertahankan dari
teks_lengkap.csv yang ada (sel kosong diisi catatan otomatis, bila ada).
Kolom jalur_disarankan dihitung ulang, kecuali bila isinya diawali
"manual:" (isian tangan yang dikunci).

Tautan akses terbuka (OA) diambil dari topik/enrich.jsonl (OpenAlex,
Semantic Scholar, Unpaywall; 28 September 2026), dari log unduhan
(unduhan/log_unduh.csv), dan dari daftar OA_MANUAL di bawah (salinan legal
yang ditemukan lewat penelusuran web, 29 September 2026: arXiv, CVF Open
Access, repositori penulis/institusi, jurnal akses terbuka). Skrip ini tidak
mengunduh apa pun dan tidak memakai jaringan.

    python3 tools/scopus/verifikasi_teks_lengkap.py
"""
import csv
import json
import re
import sys
from collections import OrderedDict
from pathlib import Path

ROOT = Path(__file__).resolve().parents[2]
KORPUS = ROOT / "literature" / "scopus-2026-09"
MATRIKS = KORPUS / "topik" / "bukti" / "matriks_bukti.csv"
KODE_C3 = KORPUS / "topik" / "bukti" / "kode_C3.txt"
ENRICH = KORPUS / "topik" / "enrich.jsonl"
LOG_UNDUH = KORPUS / "unduhan" / "log_unduh.csv"
PDF_DIR = KORPUS / "pdf"
BODY = ROOT / "manuscript" / "source" / "main6-body.tex"
OUT_CSV = KORPUS / "verifikasi" / "cek-4-teks-lengkap" / "teks_lengkap.csv"
OUT_MINTA = KORPUS / "verifikasi" / "cek-4-teks-lengkap" / "PERMINTAAN-PDF.md"
OUT_XLSX = KORPUS / "verifikasi" / "cek-4-teks-lengkap" / "PERMINTAAN-PDF-BU-FATMA.xlsx"
OUT_PRIOR = KORPUS / "unduhan" / "PRIORITAS-UNDUH-MANUAL.md"

KOLOM = ["prioritas", "kelompok", "key", "kode", "tahun", "penulis_pertama",
         "judul", "sumber", "doi", "penerbit", "pdf_ada", "jalur_disarankan",
         "email_penulis_korespondensi", "status", "tanggal_minta",
         "sumber_pdf", "catatan"]
KOLOM_TANGAN = ["email_penulis_korespondensi", "status", "tanggal_minta",
                "sumber_pdf", "catatan"]

# Angka rencana supervisor (total, dengan PDF). Selisih dilaporkan, tidak dipaksa.
RENCANA = {1: (21, 4), 2: (42, 1), 3: (124, 52), 4: (18, 6)}

# Salinan akses terbuka legal yang tidak tercatat di enrich.jsonl,
# ditemukan lewat penelusuran web 29 September 2026. Hanya arXiv, CVF Open
# Access, repositori penulis/institusi, dan jurnal akses terbuka.
OA_MANUAL = {
    "liu2018robust": "https://arxiv.org/abs/1804.00307",
    "qureshi2023seeing": "https://arxiv.org/abs/2308.07512",
    "lei2025spatio": "https://arxiv.org/abs/2409.19786",
    "meyer2024fruitnerf": "https://arxiv.org/abs/2408.06190",
    "meyer2025fruitnerf": "https://arxiv.org/abs/2505.19863",
    "zaenker2023graph": "https://arxiv.org/abs/2303.03048",
    "yang2025safe": "https://arxiv.org/abs/2409.18293",
    "santos2024multiple": "https://arxiv.org/abs/2312.16724",
    "nellithimaru2019rols": ("https://openaccess.thecvf.com/content_CVPRW_2019/papers/CVPPP/"
                             "Nellithimaru_ROLS__Robust_Object-Level_SLAM_for_Grape_Counting_CVPRW_2019_paper.pdf"),
    "matos2024tracking": ("https://openaccess.thecvf.com/content/CVPR2024W/Vision4Ag/papers/"
                          "Matos_Tracking_and_Counting_Apples_in_Orchards_Under_Intermittent_Occlusions_and_CVPRW_2024_paper.pdf"),
    "nuske2014automated": "https://www.ri.cmu.edu/pub_files/2014/9/rob21541.pdf",
    "prasetyo2020automatic": "https://gigvvy.com/journals/ijase/articles/ijase-202005-17-2-121",
    # Ditemukan pada pemeriksaan ulang 4 Oktober 2026 (OpenAlex menurut judul).
    "tan2025appleyolo": "https://doi.org/10.2139/ssrn.4861527",
}
# Tautan OA yang dicoba di peramban pada 4 Oktober 2026 dan tidak memberi PDF;
# kajian ini kembali ke jalur permintaan biasa.
OA_GAGAL = {"zheng2023object", "song2014automatic", "xieli2025pinesort", "si2026citrus"}
# Catatan versi untuk OA_MANUAL yang bukan versi terbit.
OA_CATATAN = {
    "liu2018robust": "OA = pracetak arXiv; periksa selisih dengan versi IROS",
    "qureshi2023seeing": "OA = pracetak arXiv; periksa selisih dengan versi IROS",
    "lei2025spatio": "OA = pracetak arXiv; periksa selisih dengan versi RA-L",
    "meyer2024fruitnerf": "OA = pracetak arXiv; periksa selisih dengan versi IROS",
    "meyer2025fruitnerf": "OA = pracetak arXiv; periksa selisih dengan versi IROS",
    "zaenker2023graph": "OA = pracetak arXiv; periksa selisih dengan versi IROS",
    "yang2025safe": "OA = pracetak arXiv; periksa selisih dengan versi ICRA",
    "santos2024multiple": "OA = pracetak arXiv; periksa selisih dengan versi COMPAG",
    "nellithimaru2019rols": "OA = versi CVF Open Access (setara versi terbit)",
    "matos2024tracking": "OA = versi CVF Open Access (setara versi terbit)",
    "nuske2014automated": "OA = salinan di situs Robotics Institute CMU",
    "tan2025appleyolo": "OA = pracetak SSRN; periksa selisih dengan versi Expert Systems with Applications",
}

PENERBIT = {
    "10.1016": "Elsevier", "10.1109": "IEEE", "10.23919": "IEEE",
    "10.1007": "Springer", "10.1038": "Springer Nature", "10.1002": "Wiley",
    "10.1145": "ACM", "10.1063": "AIP Publishing", "10.1088": "IOP Publishing",
    "10.1093": "Oxford University Press", "10.1017": "Cambridge University Press",
    "10.1117": "SPIE", "10.1080": "Taylor & Francis",
    "10.1051": "EDP Sciences", "10.11591": "IAES", "10.32985": "IJECES (Univ. Osijek)",
    "10.11975": "CSAE (Tiongkok)", "10.3969": "CSAE (Tiongkok)",
    "10.6041": "CSAM (Tiongkok)", "10.13925": "Journal of Fruit Science (Tiongkok)",
    "10.13031": "ASABE", "10.1364": "Optica", "10.1371": "PLOS",
    "10.15586": "Codon Publications", "10.17660": "ISHS", "10.20965": "Fuji Technology Press",
    "10.21894": "MPOB", "10.22456": "UFRGS", "10.3389": "Frontiers", "10.3390": "MDPI",
    "10.34133": "AAAS (Science Partner Journals)", "10.3920": "Wageningen Academic",
    "10.6703": "IJASE (Chaoyang Univ.)",
}
# DOI yang kosong di rekaman Scopus, ditemukan lewat penelusuran web 5 Oktober 2026.
DOI_MANUAL = {"gongal2014identification": "10.13031/aim.20141888882"}
# Jurnal yang tidak menerbitkan DOI.
PENERBIT_TANPA_DOI = {"Iaeng International Journal of Computer Science": "IAENG",
                      "Journal of Food Agriculture and Environment": "WFL Publisher"}
# Jurnal akses terbuka penuh: halaman DOI penerbit lebih andal daripada repositori.
JURNAL_OA = {"IEEE Access", "Data in Brief", "Heliyon", "Smart Agricultural Technology",
             "Plant Phenomics", "OSA Continuum", "Information Processing in Agriculture",
             "E3s Web of Conferences", "Journal of Robotics and Mechatronics"}
TIONGKOK = {"10.11975", "10.3969", "10.6041", "10.13925"}
BESAR = {"Elsevier", "IEEE", "Springer", "Springer Nature", "Wiley", "ACM",
         "AIP Publishing", "IOP Publishing", "Oxford University Press",
         "Cambridge University Press", "SPIE", "Taylor & Francis"}

JALUR_ADA = "- (PDF sudah ada)"
JALUR_LIB = "perpustakaan ULM"
JALUR_FATMA = "Bu Fatma"
JALUR_PENULIS = "penulis"


def cites(teks):
    return [k.strip() for m in re.findall(r"\\cite[a-zA-Z]*\*?(?:\[[^\]]*\])*\{([^}]*)\}", teks)
            for k in m.split(",") if k.strip()]


def potong(body, awal, akhir):
    i = body.index(awal)
    return body[i:body.index(akhir, i)]


def baca_c3():
    c3 = {}
    for baris in KODE_C3.read_text(encoding="utf-8").splitlines():
        if not baris.strip() or baris.startswith("#"):
            continue
        p = baris.split("\t")
        c3[p[0]] = {"tugas": p[1], "multipandang": p[4]}
    return c3


def url_oa(key, rec, en, log):
    """Pilih satu tautan OA terbaik, atau '' bila tidak ada."""
    if key in OA_GAGAL:
        return ""
    if key in OA_MANUAL:
        return OA_MANUAL[key]
    kand = list(en.get("pdf_urls") or []) + list(en.get("oa_landing") or [])
    note = log.get("note", "") if log else ""
    for u in re.findall(r"(https?://\S+?):\s", note + " "):
        if not u.startswith("https://doi.org/"):
            kand.append(u)
    if log and log.get("source_url"):
        kand.append(log["source_url"])
    kand = [u for u in kand if "doaj.org" not in u]
    oa_status = en.get("oa_status") or ""
    ada_oa = oa_status in ("gold", "hybrid", "bronze", "green") or bool(en.get("pdf_urls"))
    if not kand:
        return ""
    if not ada_oa and all(u.startswith("https://doi.org/") for u in kand):
        return ""

    def norm(u):
        m = re.match(r"https?://www\.ncbi\.nlm\.nih\.gov/pmc/articles/(?:PMC)?(\d+)", u)
        if m:
            return f"https://pmc.ncbi.nlm.nih.gov/articles/PMC{m.group(1)}/"
        return u

    jurnal_oa = rec.get("source", "") in JURNAL_OA

    def peringkat(u):
        if "ncbi.nlm.nih.gov" in u:
            return 0
        if "arxiv.org" in u or "thecvf.com" in u:
            return 1
        if u.startswith("https://doi.org/"):
            return 1.5 if jurnal_oa else 4
        if any(h in u for h in ("hdl.handle.net", "upcommons", "figshare", "research.wur.nl",
                                "publica.fraunhofer", "digitalcommons", "kerwa.ucr",
                                "psasir.upm", "ri.cmu.edu", "seer.ufrgs")):
            return 2
        return 3
    kand = sorted(dict.fromkeys(norm(u) for u in kand), key=peringkat)
    return kand[0]


def penerbit_dari(doi, sumber):
    if not doi:
        if "Agricultural and Biological Engineers" in sumber:
            return "ASABE"
        return PENERBIT_TANPA_DOI.get(sumber, "?")
    return PENERBIT.get(doi.split("/")[0], sumber or "?")


def jalur_untuk(r, oa):
    if r["pdf_ada"] == "Y":
        return JALUR_ADA
    if oa:
        return "OA: " + oa
    pre = r["doi"].split("/")[0] if r["doi"] else ""
    jenis = r["_doc_type"]
    prosiding = jenis in ("Conference Paper", "Book Chapter") or not r["doi"]
    if r["kode"] == "C3" and prosiding:
        return JALUR_PENULIS          # penulis Indonesia/Malaysia, mudah dihubungi
    if pre in TIONGKOK or prosiding:
        return JALUR_FATMA
    if r["penerbit"] in BESAR:
        return JALUR_LIB
    return JALUR_PENULIS


def susun():
    rows = list(csv.DictReader(MATRIKS.open(encoding="utf-8")))
    bykey = {r["key"]: r for r in rows}
    byidx = {r["idx"]: r for r in rows}
    pdfs = {p.stem for p in PDF_DIR.glob("*.pdf")}
    en = {}
    for baris in ENRICH.open(encoding="utf-8"):
        d = json.loads(baris)
        en[d.get("eid")] = d
    log = {r["key"]: r for r in csv.DictReader(LOG_UNDUH.open(encoding="utf-8"))}
    body = BODY.read_text(encoding="utf-8")
    acq = list(dict.fromkeys(cites(potong(body, r"\label{tab:acq}", r"\end{table*}"))))
    assoc = list(dict.fromkeys(cites(potong(body, r"\label{tab:assoc}", r"\end{table*}"))))
    palm = set(cites(potong(body, r"\label{sec:palm}", r"\label{sec:position}")))
    c3 = baca_c3()

    daftar = OrderedDict()   # key -> (prioritas, kelompok, catatan_awal)
    for k in acq:
        daftar[k] = (1, "Tabel 2 (tab:acq)", "")
    for k in assoc:
        if k not in daftar:
            daftar[k] = (1, "Tabel 3 (tab:assoc)", "")
    for r in rows:
        if r["kode"] == "C1" and r["dasar_keputusan"] == "judul" and r["key"] not in daftar:
            daftar[r["key"]] = (2, "C1 dikode dari judul", "")
    for r in rows:
        if r["kode"] == "C1" and r["key"] not in daftar:
            daftar[r["key"]] = (3, "C1 lainnya", "")
    for idx, c in c3.items():
        hit, mv = c["tugas"] == "HIT", c["multipandang"] == "Y"
        if not (hit or mv):
            continue
        r = byidx[idx]
        kel = "C3 " + "+".join(x for x, y in (("pencacahan", hit), ("multipandang", mv)) if y)
        cat = ("dikutip di seksi Oil palm" if r["key"] in palm
               else "tidak dikutip di seksi Oil palm (hanya masuk hitungan 11 multipandang)")
        if r["key"] not in daftar:
            daftar[r["key"]] = (4, kel, cat)

    peringatan = []
    for k in acq + assoc:
        if k not in bykey:
            peringatan.append(f"kunci tabel {k} tidak ada di matriks_bukti.csv")
        elif bykey[k]["kode"] != "C1":
            peringatan.append(f"kunci tabel {k} berkode {bykey[k]['kode']}, bukan C1")

    lama = {}
    if OUT_CSV.exists():
        lama = {r["key"]: r for r in csv.DictReader(OUT_CSV.open(encoding="utf-8"))}

    keluar = []
    for key, (pr, kel, cat) in daftar.items():
        m = bykey[key]
        r = {
            "prioritas": str(pr), "kelompok": kel, "key": key, "kode": m["kode"],
            "tahun": m["year"], "penulis_pertama": m["first_author"], "judul": m["title"],
            "sumber": m["source"], "doi": m["doi"] or DOI_MANUAL.get(key, ""),
            "penerbit": penerbit_dari(m["doi"] or DOI_MANUAL.get(key, ""), m["source"]),
            "pdf_ada": "Y" if key in pdfs else "N",
            "_doc_type": m["doc_type"],
        }
        oa = url_oa(key, m, en.get(m["eid"], {}), log.get(key))
        r["_oa"] = oa
        r["jalur_disarankan"] = jalur_untuk(r, oa)
        catatan = [c for c in (cat, OA_CATATAN.get(key, "") if oa else "") if c]
        if r["pdf_ada"] == "N" and oa.startswith("https://doi.org/"):
            catatan.append("OA menurut OpenAlex/Unpaywall; bila halaman terkunci, "
                           "pakai perpustakaan ULM")
        if key in OA_GAGAL:
            catatan.append("tautan OA dicoba di peramban 4 Oktober 2026, tidak memberi PDF")
        if m["dasar_keputusan"] == "judul":
            catatan.append("dikode dari judul")
        r["catatan"] = "; ".join(catatan)
        for kol in ("email_penulis_korespondensi", "status", "tanggal_minta", "sumber_pdf"):
            r[kol] = ""
        if key in lama:
            for kol in KOLOM_TANGAN:
                if lama[key].get(kol, "").strip():
                    r[kol] = lama[key][kol]
            if lama[key].get("jalur_disarankan", "").startswith("manual:"):
                r["jalur_disarankan"] = lama[key]["jalur_disarankan"]
        keluar.append(r)

    urut_jalur = {"OA": 0, JALUR_LIB: 1, JALUR_FATMA: 2, JALUR_PENULIS: 3}

    def kunci(r):
        j = r["jalur_disarankan"]
        jj = 0 if j.startswith("OA") else urut_jalur.get(j, 4)
        tab = 0 if r["prioritas"] == "1" else 1   # prioritas 1 tetap urut tabel naskah
        return (int(r["prioritas"]), r["pdf_ada"] == "Y",
                0 if tab == 0 else jj, 0 if tab == 0 else -int(r["tahun"] or 0))
    keluar.sort(key=kunci)
    return keluar, peringatan


def tulis_csv(keluar):
    OUT_CSV.parent.mkdir(parents=True, exist_ok=True)
    with OUT_CSV.open("w", encoding="utf-8", newline="") as f:
        w = csv.DictWriter(f, fieldnames=KOLOM, extrasaction="ignore", lineterminator="\n")
        w.writeheader()
        w.writerows(keluar)


def md_sel(t):
    return (t or "").replace("|", "\\|").replace("\n", " ")


def doi_link(doi):
    return f"[{doi}](https://doi.org/{doi})" if doi else "(tanpa DOI)"


def rekap(keluar):
    out = {}
    for p in (1, 2, 3, 4):
        g = [r for r in keluar if r["prioritas"] == str(p)]
        out[p] = (len(g), sum(r["pdf_ada"] == "Y" for r in g),
                  sum(r["pdf_ada"] == "N" and r["jalur_disarankan"].startswith("OA") for r in g))
    return out


NAMA_P = {1: "Tabel 2 dan 3 naskah", 2: "C1 dikode dari judul", 3: "C1 lainnya",
          4: "C3 sawit: pencacahan dan multipandang"}

def urut_minta(keluar):
    """Kajian tanpa PDF, diurutkan menurut prioritas, penerbit, lalu key."""
    return sorted((r for r in keluar if r["pdf_ada"] == "N"),
                  key=lambda r: (int(r["prioritas"]), r["penerbit"], r["key"]))


def tulis_permintaan(keluar):
    belum = urut_minta(keluar)
    rk = rekap(keluar)
    L = []
    a = L.append
    a("# Cek 4: Daftar Permintaan PDF Teks Lengkap")
    a("")
    a(f"Daftar ini memuat {len(belum)} kajian yang belum memiliki PDF. Isinya sama dengan "
      f"`{OUT_XLSX.name}`. Kedua berkas dihasilkan `tools/scopus/verifikasi_teks_lengkap.py` dari "
      "`teks_lengkap.csv`, sehingga tidak disunting dengan tangan. PDF yang diterima disimpan di "
      "`literature/scopus-2026-09/pdf/` dengan nama pada kolom terakhir.")
    a("")
    a("| Prioritas | Kelompok | Jumlah kajian | Sudah memiliki PDF | Belum memiliki PDF |")
    a("|---|---|---|---|---|")
    for p in (1, 2, 3, 4):
        t, y, _ = rk[p]
        a(f"| {p} | {NAMA_P[p]} | {t} | {y} | {t - y} |")
    a(f"| | **Jumlah** | {len(keluar)} | {len(keluar) - len(belum)} | {len(belum)} |")
    a("")
    a("| No. | Prioritas | Kode | Tahun | Penulis pertama | Judul | Sumber | Penerbit | DOI | Nama berkas PDF |")
    a("|---|---|---|---|---|---|---|---|---|---|")
    for i, r in enumerate(belum, 1):
        a(f"| {i} | {r['prioritas']} | {r['kode']} | {r['tahun']} | {md_sel(r['penulis_pertama'])} | "
          f"{md_sel(r['judul'])} | {md_sel(r['sumber'])} | {r['penerbit']} | {doi_link(r['doi'])} | "
          f"`{r['key']}.pdf` |")
    a("")
    OUT_MINTA.write_text("\n".join(L), encoding="utf-8")


def tulis_xlsx(keluar):
    """Daftar yang sama dengan PERMINTAAN-PDF.md dalam satu lembar Excel."""
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    except ImportError:
        print(f"openpyxl tidak tersedia; {OUT_XLSX.name} tidak ditulis", file=sys.stderr)
        return
    wb = Workbook()
    ws = wb.active
    ws.title = "Permintaan PDF"
    ws.append(["No.", "Prioritas", "Kelompok", "Kode", "Tahun", "Penulis pertama", "Judul", "Sumber",
               "Penerbit", "DOI", "Nama berkas PDF", "Status"])
    for i, r in enumerate(urut_minta(keluar), 1):
        kelompok = re.sub(r"\s*\(tab:[^)]*\)", " naskah", r["kelompok"])
        ws.append([i, int(r["prioritas"]), kelompok, r["kode"], int(r["tahun"]), r["penulis_pertama"],
                   r["judul"], r["sumber"], r["penerbit"], r["doi"], r["key"] + ".pdf", ""])
        if r["doi"]:
            c = ws.cell(row=i + 1, column=10)
            c.hyperlink = "https://doi.org/" + r["doi"]
            c.font = Font(color="0563C1", underline="single")
    tepi = Side(style="thin", color="BFBFBF")
    for baris in ws.iter_rows():
        for c in baris:
            c.border = Border(left=tepi, right=tepi, top=tepi, bottom=tepi)
            c.alignment = Alignment(vertical="top", wrap_text=True)
    for c in ws[1]:
        c.font = Font(bold=True)
        c.fill = PatternFill("solid", fgColor="D9E1F2")
        c.alignment = Alignment(vertical="center", wrap_text=True)
    for kol, lebar in zip("ABCDEFGHIJKL", [5, 9, 22, 6, 7, 18, 60, 38, 14, 30, 30, 16]):
        ws.column_dimensions[kol].width = lebar
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    wb.save(OUT_XLSX)


def tulis_prioritas(keluar):
    belum = [r for r in keluar if r["pdf_ada"] == "N"]
    rk = rekap(keluar)
    L = []
    a = L.append
    a("# Prioritas unduh manual PDF `main6`")
    a("")
    a("Dibuat oleh `tools/scopus/verifikasi_teks_lengkap.py` dari "
      "`literature/scopus-2026-09/verifikasi/cek-4-teks-lengkap/teks_lengkap.csv`; jalankan ulang skrip setelah "
      "menambah PDF (baris yang PDF-nya sudah ada akan hilang dari tabel).")
    a("")
    a("Urutan mengikuti rencana verifikasi teks lengkap: (1) kajian Tabel 2 dan 3 naskah, "
      "(2) kajian C1 yang dikode dari judul, (3) sisa kajian C1, (4) kajian sawit C3 "
      "pencacahan dan multipandang. Di dalam prioritas 2–4: OA dulu, lalu perpustakaan, "
      "Bu Fatma, penulis; tahun terbaru dulu.")
    a("")
    a("Simpan setiap PDF sebagai `literature/scopus-2026-09/pdf/<key>.pdf` (nama persis kolom "
      "*key*). Centang kotak setelah selesai. Daftar permintaan: "
      "`literature/scopus-2026-09/verifikasi/cek-4-teks-lengkap/PERMINTAAN-PDF.md`.")
    a("")
    a("| Prioritas | Kelompok | Kajian | PDF ada | Belum | OA ditemukan |")
    a("|---|---|---|---|---|---|")
    for p in (1, 2, 3, 4):
        t, y, o = rk[p]
        a(f"| {p} | {NAMA_P[p]} | {t} | {y} | {t - y} | {o} |")
    a("")
    a("| ✓ | # | Prioritas | Kode | Key (nama berkas) | Tahun | Judul | Penerbit | DOI | Jalur disarankan |")
    a("|---|---|---|---|---|---|---|---|---|---|")
    for i, r in enumerate(belum, 1):
        j = r["jalur_disarankan"]
        if j.startswith("OA: "):
            j = f"OA: <{j[4:]}>"
        a(f"| ☐ | {i} | {r['prioritas']} | {r['kode']} | `{r['key']}` | {r['tahun']} | "
          f"{md_sel(r['judul'])} | {r['penerbit']} | {doi_link(r['doi'])} | {j} |")
    a("")
    OUT_PRIOR.write_text("\n".join(L), encoding="utf-8")


def main():
    keluar, peringatan = susun()
    tulis_csv(keluar)
    tulis_permintaan(keluar)
    tulis_xlsx(keluar)
    tulis_prioritas(keluar)
    rk = rekap(keluar)
    print(f"{OUT_CSV.relative_to(ROOT)}: {len(keluar)} baris")
    for p in (1, 2, 3, 4):
        t, y, o = rk[p]
        rt, ry = RENCANA[p]
        tanda = "" if (t, y) == (rt, ry) else f"  <-- rencana {rt}/{ry}"
        print(f"  prioritas {p} ({NAMA_P[p]}): {t} kajian, {y} ber-PDF, {o} OA ditemukan{tanda}")
    for w in peringatan:
        print("PERINGATAN:", w, file=sys.stderr)


if __name__ == "__main__":
    main()
