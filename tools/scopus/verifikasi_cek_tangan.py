#!/usr/bin/env python3
"""Menyiapkan lembar Cek 2 (pemeriksaan tangan keputusan terpenting) korpus main6.

Skrip ini HANYA menyiapkan bahan. Kolom konfirmasi dibiarkan kosong dan diisi
oleh peninjau manusia; skrip tidak pernah mengisinya.

Kelompok yang disiapkan:
  cek_C1.csv        kajian C1 (menggabungkan beberapa pengamatan buah yang sama)
  cek_C3.csv        kajian C3 (pencitraan TBS kelapa sawit) beserta kodenya
  cek_eksklusi.csv  rekaman yang dikeluarkan pada tahap kelayakan (X E1-E7)
Tiap CSV juga ditulis sebagai .xlsx (daftar pilihan dan teks terbungkus) bila
openpyxl tersedia, ditambah CEK-TANGAN.md (petunjuk dan kemajuan).

Sumber (hanya dibaca, tidak diubah):
  topik/penyaringan/abstrak_*.txt   keputusan tahap kelayakan (baris `idx KODE [E#]`).
      Urutan prioritas sama dengan kode_bukti.py: berkas dibaca menurut nama
      terurut (abstrak_00..08, B00..B03, T00) dan baris terakhir menang.
      Dasar keputusan "judul" = baris di abstrak_T00.txt atau bertanda "(judul)".
  topik/penyaringan/kandidat_abstrak.csv   judul, tahun, DOI, kunci
  topik/enrich.jsonl                        abstrak
  topik/bukti/mekanisme_C1.txt, kode_C3.txt kode manual
  topik/bukti/matriks_bukti.csv             hanya untuk pemeriksaan kesesuaian
  pdf/<key>.pdf, teks/<key>.txt             teks lengkap

Aman dijalankan ulang:
  * Kolom manusia (konfirmasi, kode_benar, catatan, ...) dibaca dari CSV dan
    XLSX yang sudah ada lalu dipertahankan. Bila CSV dan XLSX berisi nilai
    berbeda untuk sel yang sama, skrip berhenti tanpa menulis apa pun.
  * Kolom *_ai dibekukan pada nilai saat baris pertama kali dibuat, agar
    konfirmasi selalu merujuk nilai yang dilihat peninjau. Bila sumber kini
    berbeda, perbedaannya ditulis di kolom catatan_skrip. Pakai --segarkan-ai
    untuk memuat ulang kolom *_ai dari sumber.
  * Baris yang tidak lagi termasuk kelompoknya tetap disimpan bila sudah diisi
    peninjau (ditandai di catatan_skrip); bila belum diisi, baris itu dibuang.
  * Nomor `no` stabil antar-eksekusi; baris baru mendapat nomor berikutnya.

Pemakaian (dari akar repo):
  python3 tools/scopus/verifikasi_cek_tangan.py
  python3 tools/scopus/verifikasi_cek_tangan.py --segarkan-ai
"""
import argparse
import csv
import datetime
import glob
import io
import json
import re
import sys
from collections import Counter
from pathlib import Path

ROOT = Path(__file__).resolve().parents[2]
KORPUS = ROOT / "literature/scopus-2026-09"
TOPIK = KORPUS / "topik"
PEN = TOPIK / "penyaringan"
BUKTI = TOPIK / "bukti"
OUT = KORPUS / "verifikasi"

KODE = re.compile(r"^(\d+)\s+(C[1-5]|T|R|M|X)(?:\s+(E\d))?(.*)$")
RENCANA = {"C1": 187, "C3": 170, "X": 153, "C1_judul": 42, "C3_judul": 35}
TENGGAT = "20 Oktober 2026"

ALASAN = {
    "E1": "bukan buah pada tanaman",
    "E2": "hanya pascapanen atau laboratorium",
    "E3": "model hasil tanpa deteksi tingkat buah",
    "E4": "hanya pemetikan atau manipulasi",
    "E5": "deteksi tanpa pencacahan atau evaluasi lain",
    "E6": "sensor non-citra",
    "E7": "bahasa selain Inggris",
}
# Petunjuk prioritas otomatis untuk lembar eksklusi (bukan keputusan).
SINYAL = [
    ("sawit", r"oil[- ]palm|palm oil|fresh fruit bunch|\bffbs?\b|palm fruit|loose fruit|elaeis"),
    ("multi-observasi", r"video|track|multi-?view|multiple views|viewpoints|both sides|two sides|"
                        r"point cloud|structure[- ]from[- ]motion|\bsfm\b|\bslam\b|re-?identification|"
                        r"double[- ]count|image sequence|frames?\b|3d reconstruction|registration"),
]

# ---------------------------------------------------------------- definisi lembar
LEMBAR = {
    "C1": {
        "berkas": "cek_C1",
        "ai": ["dasar_keputusan_ai", "mekanisme_ai", "akuisisi_ai", "per_kelas_ai", "hasil_ringkas_ai"],
        "manusia": ["konfirmasi_C1", "kode_benar", "catatan"],
        "pilihan": {"konfirmasi_C1": "Y,N", "kode_benar": "Y,N"},
    },
    "C3": {
        "berkas": "cek_C3",
        "ai": ["dasar_keputusan_ai", "tugas_ai", "lokasi_ai", "modalitas_ai", "multipandang_ai", "per_kelas_ai"],
        "manusia": ["konfirmasi_C3", "kode_benar", "catatan"],
        "pilihan": {"konfirmasi_C3": "Y,N", "kode_benar": "Y,N"},
    },
    "X": {
        "berkas": "cek_eksklusi",
        "ai": ["dasar_keputusan_ai", "alasan_ai"],
        "manusia": ["konfirmasi_alasan", "seharusnya_C1_atau_C3", "catatan"],
        "pilihan": {"konfirmasi_alasan": "Y,N", "seharusnya_C1_atau_C3": "C1,C3,C1+C3,TIDAK"},
    },
}
DEPAN = ["no", "idx", "key", "tahun", "judul", "doi", "abstrak"]
BELAKANG = ["pdf_ada", "path_teks", "hanya_judul"]


def kolom_lembar(g):
    d = LEMBAR[g]
    extra = ["sinyal_kata_kunci"] if g == "X" else []
    return DEPAN + d["ai"] + BELAKANG + extra + d["manusia"] + ["catatan_skrip"]


# ---------------------------------------------------------------- pemuat sumber
def muat_keputusan():
    dec, basis, asal, kemunculan = {}, {}, {}, Counter()
    for fn in sorted(glob.glob(str(PEN / "abstrak_*.txt"))):
        nama = Path(fn).name
        for line in open(fn, encoding="utf-8"):
            m = KODE.match(line.strip())
            if not m:
                continue
            i = int(m.group(1))
            dec[i] = (m.group(2), m.group(3) or "")
            basis[i] = "judul" if "(judul)" in line or nama == "abstrak_T00.txt" else "abstrak"
            asal[i] = nama
            kemunculan[i] += 1
    return dec, basis, asal, kemunculan


def muat_tab(path, n):
    out = {}
    for line in open(path, encoding="utf-8"):
        if line.startswith("#") or not line.strip():
            continue
        p = line.rstrip("\n").split("\t")
        p += [""] * (n - len(p))
        out[int(p[0])] = p[1:n]
    return out


def ada_teks(key):
    pdf = (KORPUS / "pdf" / f"{key}.pdf").exists()
    teks = KORPUS / "teks" / f"{key}.txt"
    return ("Y" if pdf else "N"), (teks.relative_to(ROOT).as_posix() if teks.exists() else "")


def bangun_baris():
    dec, basis, asal, kemunculan = muat_keputusan()
    kand = {int(r["idx"]): r for r in csv.DictReader(open(PEN / "kandidat_abstrak.csv", encoding="utf-8"))}
    enr = {}
    for line in open(TOPIK / "enrich.jsonl", encoding="utf-8"):
        e = json.loads(line)
        enr[e["eid"]] = e
    mek = muat_tab(BUKTI / "mekanisme_C1.txt", 5)
    c3 = muat_tab(BUKTI / "kode_C3.txt", 6)
    info = {"dec": dec, "basis": basis, "asal": asal, "kemunculan": kemunculan,
            "kand": kand, "mek": mek, "c3": c3, "peringatan": []}

    ganda = {i: n for i, n in kemunculan.items() if n > 1}
    if ganda:
        info["peringatan"].append(f"{len(ganda)} idx muncul di lebih dari satu berkas penyaringan; "
                                  f"dipakai baris terakhir menurut urutan nama berkas: {sorted(ganda)[:20]}")
    tanpa = sorted(set(kand) - set(dec))
    if tanpa:
        info["peringatan"].append(f"{len(tanpa)} kandidat tanpa keputusan kelayakan: {tanpa[:20]}")

    grup = {"C1": {}, "C3": {}, "X": {}}
    for i, (kode, alasan) in dec.items():
        if kode not in grup:
            continue
        r = kand.get(i)
        if r is None:
            info["peringatan"].append(f"idx {i} ada di berkas penyaringan tetapi tidak di kandidat_abstrak.csv")
            continue
        abstrak = (enr.get(r["eid"], {}).get("abstract") or "").strip()
        pdf, teks = ada_teks(r["key"])
        b = {"idx": i, "key": r["key"], "tahun": r["year"], "judul": r["title"], "doi": r["doi"],
             "abstrak": abstrak, "dasar_keputusan_ai": basis[i], "pdf_ada": pdf, "path_teks": teks,
             "hanya_judul": "Y" if basis[i] == "judul" else "N", "_catatan": []}
        if basis[i] == "judul" and abstrak:
            b["_catatan"].append("diputuskan dari judul, tetapi abstrak kini tersedia di enrich.jsonl")
        if basis[i] == "judul" and not abstrak and pdf == "N":
            b["_catatan"].append("tanpa abstrak dan tanpa PDF: cari abstrak/teks lengkap lewat DOI")
        if kode == "C1":
            m = mek.get(i)
            if m is None:
                b["_catatan"].append("tidak ada baris di mekanisme_C1.txt")
                m = ["", "", "", ""]
            b.update(mekanisme_ai=m[0], akuisisi_ai=m[1], per_kelas_ai=m[2], hasil_ringkas_ai=m[3])
        elif kode == "C3":
            k = c3.get(i)
            if k is None:
                b["_catatan"].append("tidak ada baris di kode_C3.txt")
                k = ["", "", "", "", ""]
            b.update(tugas_ai=k[0], lokasi_ai=k[1], modalitas_ai=k[2], multipandang_ai=k[3], per_kelas_ai=k[4])
            tanda = [n for n, v in zip(["tugas", "lokasi", "modalitas"], k[:3]) if v == "?"]
            if tanda:
                b["_catatan"].append("kode '?' pada: " + ", ".join(tanda))
        else:
            b["alasan_ai"] = f"{alasan} ({ALASAN.get(alasan, '?')})" if alasan else ""
            teks_s = f"{r['title']} {abstrak}"
            b["sinyal_kata_kunci"] = ";".join(n for n, p in SINYAL if re.search(p, teks_s, re.I))
        grup[kode][i] = b

    # pemeriksaan kesesuaian dengan matriks_bukti.csv
    mat_path = BUKTI / "matriks_bukti.csv"
    if mat_path.exists():
        mat = {int(r["idx"]): r for r in csv.DictReader(open(mat_path, encoding="utf-8"))}
        for g in ("C1", "C3"):
            a = {i for i, r in mat.items() if r["kode"] == g}
            if a != set(grup[g]):
                info["peringatan"].append(f"anggota {g} di matriks_bukti.csv berbeda dari berkas penyaringan "
                                          f"(jalankan kode_bukti.py): +{sorted(set(grup[g]) - a)[:10]} "
                                          f"-{sorted(a - set(grup[g]))[:10]}")
        beda = [i for i, b in grup["C1"].items() if i in mat and
                (mat[i]["mekanisme"], mat[i]["akuisisi_manual"], mat[i]["per_kelas"], mat[i]["hasil_ringkas"])
                != (b["mekanisme_ai"], b["akuisisi_ai"], b["per_kelas_ai"], b["hasil_ringkas_ai"])]
        if beda:
            info["peringatan"].append(f"{len(beda)} kode C1 di matriks_bukti.csv berbeda dari mekanisme_C1.txt "
                                      f"(jalankan kode_bukti.py): {beda[:10]}")
    lebih = sorted(set(mek) - set(grup["C1"]))
    if lebih:
        info["peringatan"].append(f"mekanisme_C1.txt memuat idx yang bukan C1: {lebih[:20]}")
    lebih = sorted(set(c3) - set(grup["C3"]))
    if lebih:
        info["peringatan"].append(f"kode_C3.txt memuat idx yang bukan C3: {lebih[:20]}")
    return grup, info


# ---------------------------------------------------------------- baca isian lama
def baca_csv_lama(path):
    if not path.exists():
        return {}
    raw = path.read_bytes()
    for enc in ("utf-8-sig", "cp1252"):
        try:
            teks = raw.decode(enc)
            break
        except UnicodeDecodeError:
            continue
    baris_pertama = teks.splitlines()[0] if teks else ""
    delim = ";" if baris_pertama.count(";") > baris_pertama.count(",") else ","
    out = {}
    for r in csv.DictReader(io.StringIO(teks, newline=""), delimiter=delim):
        if r.get("idx", "").strip():
            out[int(float(r["idx"]))] = {k: (v or "").strip() for k, v in r.items() if k}
    return out


def baca_xlsx_lama(path):
    if not path.exists():
        return {}
    import openpyxl
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb["cek"] if "cek" in wb.sheetnames else wb.worksheets[0]
    rows = ws.iter_rows(values_only=True)
    hdr = [str(h) if h is not None else "" for h in next(rows)]
    out = {}
    for row in rows:
        r = {h: ("" if v is None else str(v).strip()) for h, v in zip(hdr, row) if h}
        if r.get("idx"):
            out[int(float(r["idx"]))] = r
    return out


def gabung_lama(g, pakai_xlsx):
    d = LEMBAR[g]
    lama_csv = baca_csv_lama(OUT / f"{d['berkas']}.csv")
    lama_xlsx = baca_xlsx_lama(OUT / f"{d['berkas']}.xlsx") if pakai_xlsx else {}
    konflik, gabung = [], {}
    for i in set(lama_csv) | set(lama_xlsx):
        a, b = lama_csv.get(i, {}), lama_xlsx.get(i, {})
        r = dict(a or b)
        for k in d["manusia"]:
            va, vb = a.get(k, ""), b.get(k, "")
            if va and vb and va != vb:
                konflik.append(f"{d['berkas']} idx {i} kolom {k}: CSV={va!r} XLSX={vb!r}")
            r[k] = va or vb
        gabung[i] = r
    return gabung, konflik


# ---------------------------------------------------------------- tulis
def tulis_csv(path, kolom, baris):
    with open(path, "w", newline="", encoding="utf-8-sig") as f:
        w = csv.DictWriter(f, fieldnames=kolom, extrasaction="ignore")
        w.writeheader()
        w.writerows(baris)


def tulis_xlsx(path, g, kolom, baris):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.datavalidation import DataValidation

    d = LEMBAR[g]
    wb = Workbook()
    ws = wb.active
    ws.title = "cek"
    ws.append(kolom)
    for b in baris:
        ws.append([int(b[k]) if k in ("no", "idx") and str(b.get(k, "")).strip() else b.get(k, "") for k in kolom])
    lebar = {"no": 5, "idx": 7, "key": 24, "tahun": 7, "judul": 45, "doi": 22, "abstrak": 80,
             "hasil_ringkas_ai": 40, "alasan_ai": 30, "path_teks": 28, "catatan": 40,
             "catatan_skrip": 36, "sinyal_kata_kunci": 18, "dasar_keputusan_ai": 11}
    bungkus = {"judul", "abstrak", "hasil_ringkas_ai", "alasan_ai", "catatan", "catatan_skrip"}
    kuning = PatternFill("solid", fgColor="FFF2CC")
    abu = PatternFill("solid", fgColor="D9D9D9")
    n = len(baris) + 1
    for j, k in enumerate(kolom, 1):
        huruf = get_column_letter(j)
        ws.column_dimensions[huruf].width = lebar.get(k, 13)
        h = ws.cell(row=1, column=j)
        h.font = Font(bold=True)
        h.fill = kuning if k in d["manusia"] else abu
        h.alignment = Alignment(wrap_text=True, vertical="top")
        for i in range(2, n + 1):
            c = ws.cell(row=i, column=j)
            c.alignment = Alignment(wrap_text=k in bungkus, vertical="top")
            if k in d["manusia"]:
                c.fill = kuning
    for i in range(2, n + 1):
        ws.row_dimensions[i].height = 110
    for k, daftar in d["pilihan"].items():
        dv = DataValidation(type="list", formula1=f'"{daftar}"', allow_blank=True,
                            showErrorMessage=True, errorTitle="Nilai tidak sah",
                            error=f"Pilih salah satu: {daftar}")
        huruf = get_column_letter(kolom.index(k) + 1)
        dv.add(f"{huruf}2:{huruf}{max(n, 2) + 200}")
        ws.add_data_validation(dv)
    ws.freeze_panes = "F2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(kolom))}{n}"

    ket = wb.create_sheet("keterangan")
    ket.column_dimensions["A"].width = 26
    ket.column_dimensions["B"].width = 100
    for r in keterangan(g):
        ket.append(r)
    for c in ket["A"]:
        c.font = Font(bold=True)
    wb.save(path)


def keterangan(g):
    umum = [
        ("Petunjuk", "Isi hanya kolom berlatar kuning. Kolom *_ai adalah keputusan/kode model bahasa (dibekukan)."),
        ("", "Perubahan keputusan dicatat di berkas sumber dan di verifikasi/log_perubahan.csv (lihat CEK-TANGAN.md)."),
        ("hanya_judul", "Y = keputusan dibuat dari judul/sumber saja; cari abstrak atau teks lengkap dulu (DOI)."),
        ("pdf_ada / path_teks", "PDF di literature/scopus-2026-09/pdf/<key>.pdf; teks di path_teks."),
        ("catatan_skrip", "Diisi skrip: kode '?', abstrak baru tersedia, atau sumber kini berbeda dari *_ai."),
    ]
    if g == "C1":
        return umum + [
            ("konfirmasi_C1", "Y = kajian sungguh menggabungkan beberapa pengamatan buah yang sama "
                              "(video, beberapa pandang, pindaian berulang). N = bukan; tulis kode yang benar di catatan."),
            ("kode_benar", "Y = mekanisme, akuisisi, dan per_kelas benar. N = tulis koreksi di catatan."),
            ("mekanisme", "M0 tanpa asosiasi; M1 koreksi statistik; M2 penampilan/re-ID; M3 pelacakan temporal; "
                          "M4 geometri/3D; M5 asosiasi dipelajari; DATA set data. Gabungan ditulis M3+M4."),
            ("akuisisi", "V video lintasan; D beberapa pandang diskret; S pindaian 3D/awan titik; "
                         "T kunjungan ulang antarwaktu; 1 satu pandang."),
            ("per_kelas", "Y bila hitungan dilaporkan per kelas (mis. kematangan)."),
        ]
    if g == "C3":
        return umum + [
            ("konfirmasi_C3", "Y = kajian memang mencitrakan TBS kelapa sawit (termasuk grading pabrik dan brondolan)."),
            ("kode_benar", "Y = tugas, lokasi, modalitas, multipandang, dan per_kelas benar. N = tulis koreksi di catatan."),
            ("tugas", "KLS klasifikasi kematangan citra tandan; DET deteksi tandan dalam adegan; HIT pencacahan; "
                      "LF brondolan; MUTU cacat/USB/minyak/massa/volume/pose; SENSOR modalitas non-RGB; DATA set data."),
            ("lokasi", "POHON; TANAH (tanah/TPH); PABRIK (pabrik/konveyor/tumpukan/laboratorium); UAV; ? tidak jelas."),
            ("multipandang", "Y bila memakai beberapa pandang/pengamatan tandan yang sama."),
            ("per_kelas", "Y bila kelas kematangan diberikan pada instans."),
        ]
    return umum + [
        ("konfirmasi_alasan", "Y = alasan eksklusi (E1-E7) benar. N = tulis alasan/kode yang benar di catatan."),
        ("seharusnya_C1_atau_C3", "C1, C3, C1+C3 bila rekaman ini seharusnya masuk; TIDAK bila memang bukan C1/C3 "
                                  "(kode inklusi lain, mis. C2, ditulis di catatan)."),
        ("sinyal_kata_kunci", "Petunjuk prioritas otomatis dari judul+abstrak (sawit / multi-observasi); bukan keputusan."),
    ] + [(k, v) for k, v in ALASAN.items()]


def tulis_md(stat, info, ada_xlsx):
    s = stat
    ext = "`.csv` dan `.xlsx`" if ada_xlsx else "`.csv`"
    lembar_isi = ("Isi lembar `.xlsx` (ada daftar pilihan dan teks terbungkus). Jangan mengisi CSV dan XLSX "
                  "sekaligus untuk baris yang sama; bila keduanya berbeda, skrip berhenti dan meminta Anda "
                  "menyelaraskannya. Untuk mengosongkan isian, kosongkan di kedua berkas (skrip "
                  "mengambil nilai yang tidak kosong)." if ada_xlsx else
                  "Isi lembar CSV (openpyxl tidak tersedia, sehingga XLSX tidak dibuat).")
    peringatan = "\n".join(f"- {p}" for p in info["peringatan"]) or "- Tidak ada."
    md = f"""# Cek 2 — Pemeriksaan Tangan Keputusan Terpenting

> Berkas ini dibuat ulang oleh `tools/scopus/verifikasi_cek_tangan.py` setiap kali
> skrip dijalankan (terakhir {datetime.date.today().isoformat()}). Jangan disunting
> tangan; ubah skripnya.

## Tujuan

Penyaringan dan pengodean korpus `main6` dilakukan satu peninjau dengan bantuan
model bahasa besar (`PROTOKOL.md` §6). Cek 2 memastikan keputusan yang paling
menentukan isi naskah benar-benar diperiksa manusia. **Setiap keputusan di cek
ini adalah keputusan peninjau.** Lembar hanya menampilkan keputusan dan kode
model beserta bukti untuk menilainya; kolom konfirmasi sengaja kosong dan tidak
pernah diisi skrip.

## Lembar kerja

Lembar tersedia dalam format {ext} di folder ini. {lembar_isi}

| Kelompok | Lembar | Jumlah | Rencana | Diputuskan dari judul saja | Sudah dikonfirmasi |
|---|---|---:|---:|---:|---:|
""".rstrip("\n") + "\n"
    rows = []
    for g, nama in (("C1", "Kajian inti C1"), ("C3", "Kajian sawit C3"), ("X", "Eksklusi tahap kelayakan")):
        rows.append(f"| {nama} | `{LEMBAR[g]['berkas']}` | {s[g]['n']} | {RENCANA[g]} | {s[g]['judul']} | {s[g]['isi']} |")
    md += "\n".join(rows) + "\n"
    md += f"""
Rencana kelompok "judul saja": C1 {RENCANA['C1_judul']}, C3 {RENCANA['C3_judul']}.

Kolom penting:

- `hanya_judul = Y`: keputusan model dibuat dari judul dan sumber saja (baris di
  `abstrak_T00.txt` atau bertanda `(judul)`). Kolom `abstrak` biasanya kosong.
- `pdf_ada`, `path_teks`: teks lengkap ada di `literature/scopus-2026-09/pdf/<key>.pdf`
  dan `teks/<key>.txt`.
- `catatan_skrip`: diisi skrip (kode `?`, abstrak yang baru tersedia, sumber yang
  kini berbeda dari kolom `*_ai`, baris yang tidak lagi termasuk kelompoknya).
- Kolom berlatar kuning (`konfirmasi_*`, `kode_benar`, `seharusnya_C1_atau_C3`,
  `catatan`) adalah milik peninjau. Arti kode ada di lembar `keterangan`.

## Urutan kerja yang disarankan

1. **Rekaman yang diputuskan dari judul saja** (saring `hanya_judul = Y`) di
   `cek_C1` lalu `cek_C3`. Sebelum memutuskan, cari dulu abstrak atau teks
   lengkapnya (DOI, halaman penerbit, PDF lokal bila `pdf_ada = Y`). Bila tidak
   ada yang dapat dibaca, tulis di `catatan` sumber apa yang sudah dicoba.
2. **Sisa `cek_C1`**: apakah kajian sungguh menggabungkan beberapa pengamatan
   buah yang sama (video, beberapa pandang, pindaian berulang)? Isi
   `konfirmasi_C1`; lalu periksa mekanisme, akuisisi, dan per_kelas → `kode_benar`.
3. **Sisa `cek_C3`**: apakah kajian mencitrakan TBS kelapa sawit? Isi
   `konfirmasi_C3`; lalu periksa tugas, lokasi, modalitas, dan multipandang →
   `kode_benar`. Kode `?` (lihat `catatan_skrip`) diputuskan dari teks lengkap bila ada.
4. **`cek_eksklusi`**: apakah alasan E1–E7 benar (`konfirmasi_alasan`), dan apakah
   ada kajian C1 atau C3 yang hilang (`seharusnya_C1_atau_C3`). Dahulukan baris
   dengan `hanya_judul = Y` dan baris yang `sinyal_kata_kunci`-nya berisi `sawit`
   atau `multi-observasi` (petunjuk prioritas otomatis, bukan keputusan).

## Mencatat perubahan

Lembar cek hanya mencatat penilaian. Bila penilaian mengubah keputusan atau kode,
ubah juga berkas sumbernya, lalu catat setiap perubahan di log.

| Yang berubah | Berkas yang disunting |
|---|---|
| Keputusan penyaringan (C1/C2/C3/…/X, alasan E#) | baris `idx KODE [E#]` di `topik/penyaringan/abstrak_*.txt` |
| Kode mekanisme, akuisisi, per_kelas, hasil ringkas C1 | `topik/bukti/mekanisme_C1.txt` |
| Kode tugas, lokasi, modalitas, multipandang, per_kelas C3 | `topik/bukti/kode_C3.txt` |
| **Setiap** perubahan di atas | `verifikasi/log_perubahan.csv` |

Aturan penyuntingan berkas penyaringan:

- Saat ini setiap idx muncul tepat satu kali di seluruh `abstrak_*.txt`. Sunting
  baris itu di tempatnya; jangan menambah baris kedua di berkas lain. Bila suatu
  idx sampai muncul lebih dari sekali, `kode_bukti.py` memakai baris terakhir
  menurut urutan nama berkas (`abstrak_00` … `abstrak_08`, `B00` … `B03`, `T00`),
  sedangkan hitungan "judul saja" di `gambar_tinjauan.py` menghitung setiap baris,
  sehingga angka PRISMA menjadi salah.
- Dasar keputusan "judul" ditentukan oleh letak baris di `abstrak_T00.txt` atau
  tanda `(judul)`. Bila Anda memutuskan ulang setelah membaca abstrak, catat
  `kolom = dasar_keputusan` di log; memindahkan baris antarberkas mengubah angka
  254 di `PROTOKOL.md` dan naskah, jadi putuskan itu secara sadar.
- Rekaman yang menjadi C1 perlu baris baru di `mekanisme_C1.txt`; yang menjadi C3
  perlu baris baru di `kode_C3.txt`. Rekaman yang keluar dari C1/C3 dihapus
  barisnya dari berkas kode itu.

Format `log_perubahan.csv` (satu baris per sel yang berubah):

```
tanggal,berkas,idx_atau_key,kolom,nilai_lama,nilai_baru,alasan,oleh
2026-10-05,topik/penyaringan/abstrak_T00.txt,1234,kode,C2,C1,"abstrak: pelacakan video dengan garis hitung",MZM
```

`tanggal` ISO (YYYY-MM-DD); `berkas` relatif terhadap `literature/scopus-2026-09/`;
`kolom` misalnya `kode`, `alasan`, `mekanisme`, `akuisisi`, `tugas`, `lokasi`,
`multipandang`, `per_kelas`, `dasar_keputusan`; `oleh` inisial peninjau.

## Setelah ada perubahan

1. `python3 tools/scopus/kode_bukti.py` (matriks bukti), lalu
   `gambar_tinjauan.py` dan `tabel_lampiran.py`.
2. Perbarui angka di naskah `main6`, `PROTOKOL.md`, dan `AGENTS.md` §3
   (`AGENTS.md` §7).
3. Jalankan ulang skrip ini. Isian peninjau dipertahankan; kolom `*_ai` tetap
   menunjukkan nilai yang dinilai, dan `catatan_skrip` mencatat nilai sumber yang
   kini berbeda.

## Tenggat

**{TENGGAT}**: kirim `log_perubahan.csv` sejauh yang sudah terisi kepada Fatma,
meskipun Cek 2 belum selesai seluruhnya.

## Peringatan dari eksekusi terakhir

{peringatan}
"""
    (OUT / "CEK-TANGAN.md").write_text(md, encoding="utf-8")


# ---------------------------------------------------------------- utama
def main():
    ap = argparse.ArgumentParser(description=__doc__.split("\n")[0])
    ap.add_argument("--segarkan-ai", action="store_true",
                    help="muat ulang kolom *_ai dari sumber (isian peninjau tetap dipertahankan)")
    args = ap.parse_args()

    try:
        import openpyxl  # noqa: F401
        ada_xlsx = True
    except ImportError:
        ada_xlsx = False
        print("openpyxl tidak tersedia: XLSX dilewati, isian XLSX lama tidak dibaca.", file=sys.stderr)

    OUT.mkdir(parents=True, exist_ok=True)
    grup, info = bangun_baris()

    # baca semua isian lama lebih dulu; berhenti bila ada konflik
    lama, konflik = {}, []
    for g in LEMBAR:
        lama[g], k = gabung_lama(g, ada_xlsx)
        konflik += k
    if konflik:
        print("BERHENTI: isian CSV dan XLSX berbeda. Selaraskan dulu, tidak ada berkas yang ditulis.", file=sys.stderr)
        for k in konflik:
            print("  " + k, file=sys.stderr)
        sys.exit(1)

    stat = {}
    siap = {}
    for g, d in LEMBAR.items():
        kolom = kolom_lembar(g)
        kini, sebelum = grup[g], lama[g]
        no_maks = max([int(float(r["no"])) for r in sebelum.values() if r.get("no", "").strip()] or [0])
        hasil = []
        for i in sorted(set(kini) | set(sebelum)):
            b, r = kini.get(i), sebelum.get(i)
            terisi = r is not None and any(r.get(k, "") for k in d["manusia"])
            if b is None and not terisi:
                continue  # keluar dari kelompok dan belum dinilai: buang
            if b is None:
                baris = {k: r.get(k, "") for k in kolom}
                kode, al = info["dec"].get(i, ("?", ""))
                baris["catatan_skrip"] = f"tidak lagi {g} di sumber (kini {kode} {al}".rstrip() + \
                    "); baris dipertahankan karena sudah diisi peninjau"
                hasil.append(baris)
                continue
            baris = {k: b.get(k, "") for k in kolom}
            catatan = list(b["_catatan"])
            if r is not None:
                baris["no"] = r.get("no", "")
                for k in d["manusia"]:
                    baris[k] = r.get(k, "")
                if not args.segarkan_ai:
                    beda = [f"{k}={b.get(k, '')!r}" for k in d["ai"] if r.get(k, "") != b.get(k, "")]
                    for k in d["ai"]:
                        baris[k] = r.get(k, "")
                    if beda:
                        catatan.append("sumber kini: " + "; ".join(beda))
            baris["catatan_skrip"] = " | ".join(catatan)
            hasil.append(baris)
        # nomor stabil: baris lama mempertahankan no; baris baru diurutkan (judul-saja dulu)
        baru = [b for b in hasil if not str(b.get("no", "")).strip()]
        if g == "X":
            baru.sort(key=lambda b: (b["hanya_judul"] != "Y", b["alasan_ai"], int(b["idx"])))
        else:
            baru.sort(key=lambda b: (b["hanya_judul"] != "Y", -int(b["tahun"] or 0), b["key"]))
        for b in baru:
            no_maks += 1
            b["no"] = no_maks
        hasil.sort(key=lambda b: int(float(b["no"])))
        siap[g] = (kolom, hasil)
        stat[g] = {"n": len(kini), "judul": sum(1 for b in kini.values() if b["hanya_judul"] == "Y"),
                   "isi": sum(1 for b in hasil if b.get(d["manusia"][0], "")),
                   "yatim": len(hasil) - sum(1 for b in hasil if int(b["idx"]) in kini)}

    for g, (kolom, hasil) in siap.items():
        nama = LEMBAR[g]["berkas"]
        try:
            tulis_csv(OUT / f"{nama}.csv", kolom, hasil)
            if ada_xlsx:
                tulis_xlsx(OUT / f"{nama}.xlsx", g, kolom, hasil)
        except PermissionError as e:
            print(f"Gagal menulis {nama}: {e}. Tutup berkas di Excel lalu jalankan ulang.", file=sys.stderr)
            sys.exit(1)

    log = OUT / "log_perubahan.csv"
    try:
        with open(log, "x", newline="", encoding="utf-8") as f:
            f.write("tanggal,berkas,idx_atau_key,kolom,nilai_lama,nilai_baru,alasan,oleh\n")
        print("log_perubahan.csv dibuat (baru).")
    except FileExistsError:
        pass

    tulis_md(stat, info, ada_xlsx)

    # ringkasan ke layar
    for g in LEMBAR:
        s = stat[g]
        print(f"{LEMBAR[g]['berkas']:14s} n={s['n']:4d} (rencana {RENCANA[g]}) judul-saja={s['judul']:3d} "
              f"sudah-dikonfirmasi={s['isi']} baris-yatim-dipertahankan={s['yatim']}")
    print(f"rencana judul-saja: C1 {RENCANA['C1_judul']}, C3 {RENCANA['C3_judul']}")
    c1 = grup["C1"].values()
    c3 = grup["C3"].values()
    x = grup["X"].values()
    print(f"C1 tanpa abstrak={sum(1 for b in c1 if not b['abstrak'])} pdf={sum(b['pdf_ada'] == 'Y' for b in c1)} "
          f"judul-saja+pdf={sum(b['pdf_ada'] == 'Y' for b in c1 if b['hanya_judul'] == 'Y')}")
    print(f"C3 tanpa abstrak={sum(1 for b in c3 if not b['abstrak'])} pdf={sum(b['pdf_ada'] == 'Y' for b in c3)} "
          f"judul-saja+pdf={sum(b['pdf_ada'] == 'Y' for b in c3 if b['hanya_judul'] == 'Y')} "
          f"judul-saja+abstrak-kini={sum(1 for b in c3 if b['hanya_judul'] == 'Y' and b['abstrak'])}")
    print(f"X   tanpa abstrak={sum(1 for b in x if not b['abstrak'])} judul-saja={sum(b['hanya_judul'] == 'Y' for b in x)} "
          f"alasan={dict(sorted(Counter(b['alasan_ai'][:2] for b in x).items()))} "
          f"sinyal={dict(Counter(b['sinyal_kata_kunci'] for b in x if b['sinyal_kata_kunci']))}")
    for p in info["peringatan"]:
        print("PERINGATAN:", p)


if __name__ == "__main__":
    main()
